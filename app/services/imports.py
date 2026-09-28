"""Playlist import parsing and normalization."""

from __future__ import annotations

import base64
import json
import re
import zipfile
from io import BytesIO
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook

from ..list_sources import PlaylistSourceFetcher, parse_csv_items
from ..security import safe_error
from ..domain.titles import normalized_title_key
from ..util import first_value
from ..schemas import ImportPayload, MAX_IMPORT_ROWS

MAX_XLSX_UNCOMPRESSED = 200 * 1024 * 1024
MAX_XLSX_ENTRIES = 5000
MAX_JSON_PAYLOAD_BYTES = 10 * 1024 * 1024


def _check_xlsx_bomb(content: bytes) -> None:
    """Reject zip-bomb style xlsx before openpyxl decompresses it into memory."""
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ENTRIES:
                raise HTTPException(413, "xlsx 内部条目过多")
            total = sum(info.file_size for info in entries)
            if total > MAX_XLSX_UNCOMPRESSED:
                raise HTTPException(413, "xlsx 解压后数据过大")
    except HTTPException:
        raise
    except zipfile.BadZipFile as exc:
        raise HTTPException(422, "无法读取 xlsx：不是有效的压缩包") from exc


def parse_xlsx(encoded: str) -> tuple[str | None, list[dict[str, Any]]]:
    workbook = None
    try:
        content = base64.b64decode(encoded, validate=True)
        _check_xlsx_bomb(content)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, f"无法读取 xlsx：{safe_error(exc)}") from exc
    try:
        sheet = workbook.active
        first_row = next(sheet.iter_rows(min_row=1, max_row=1), None)
        if not first_row:
            raise HTTPException(422, "xlsx 为空，请确认文件包含表头和影片数据")
        headers = [str(cell.value or "").strip() for cell in first_row]
        required = {"总排名", "IMDb ID", "英文/原片名", "年份", "中文译名"}
        if not required.issubset(headers):
            raise HTTPException(422, "xlsx 缺少必需列：总排名、IMDb ID、英文/原片名、年份、中文译名")
        positions = {header: index for index, header in enumerate(headers)}
        items: list[dict[str, Any]] = []
        for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if row_index > MAX_IMPORT_ROWS:
                raise HTTPException(413, f"xlsx 行数超过上限（{MAX_IMPORT_ROWS} 行）")
            title = row[positions["英文/原片名"]]
            if not title:
                continue
            year = row[positions["年份"]]
            items.append({
                "rank_no": row[positions["总排名"]],
                "imdb_id": row[positions["IMDb ID"]],
                "original_title": str(title).strip(),
                "year": int(year) if str(year or "").isdigit() else None,
                "chinese_title": row[positions["中文译名"]],
            })
        return sheet.title, items
    finally:
        workbook.close()


def parse_json(data: dict[str, Any] | list[dict[str, Any]]) -> tuple[str | None, list[dict[str, Any]]]:
    payload_size = len(json.dumps(data, ensure_ascii=False).encode("utf-8"))
    if payload_size > MAX_JSON_PAYLOAD_BYTES:
        raise HTTPException(413, "JSON 片单数据过大")
    source = data if isinstance(data, dict) else {"films": data}
    films = source.get("films") or source.get("items") or []
    if not isinstance(films, list):
        raise HTTPException(422, "JSON 必须包含 films 或 items 数组")
    if len(films) > MAX_IMPORT_ROWS:
        raise HTTPException(413, f"JSON 条目数超过上限（{MAX_IMPORT_ROWS} 条）")
    items: list[dict[str, Any]] = []
    for index, film in enumerate(films, start=1):
        if not isinstance(film, dict):
            continue
        title = first_value(film, ("original_title", "title", "english_title", "英文/原片名"))
        if not title:
            continue
        year = first_value(film, ("year", "年份"))
        items.append({
            "rank_no": first_value(film, ("rank_no", "rank", "总排名"), index),
            "imdb_id": first_value(film, ("imdb_id", "imdbId", "imdb", "IMDb ID")),
            "original_title": str(title).strip(),
            "year": int(year) if str(year or "").isdigit() else None,
            "chinese_title": first_value(film, ("chinese_title", "cn_title", "中文译名")),
            "tmdb_id": first_value(film, ("tmdb_id", "tmdbId", "tmdb")),
        })
    return source.get("listName") or source.get("name"), items


def normalize_import_items(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in items:
        title = str(item.get("original_title") or "").strip()
        imdb_id = str(item.get("imdb_id") or "").strip() or None
        tmdb_raw = item.get("tmdb_id")
        tmdb_id = int(tmdb_raw) if str(tmdb_raw or "").isdigit() else None
        year_raw = item.get("year")
        year = int(year_raw) if str(year_raw or "").isdigit() else None
        if not title:
            continue
        normalized_title = normalized_title_key(title)
        key = f"tmdb:{tmdb_id}" if tmdb_id else (f"imdb:{imdb_id.lower()}" if imdb_id else f"title:{normalized_title}:{year or ''}")
        if key in seen:
            continue
        seen.add(key)
        source_ref = str(item.get("source_ref") or "").strip() or None
        result.append({"rank_no": len(result) + 1, "imdb_id": imdb_id, "original_title": title, "year": year,
                       "chinese_title": item.get("chinese_title"), "tmdb_id": tmdb_id, "source_ref": source_ref})
    return result


async def resolve_import(payload: ImportPayload) -> tuple[str | None, list[dict[str, Any]], dict[str, Any]]:
    if payload.source_url:
        source = await PlaylistSourceFetcher().fetch(payload.source_url, payload.limit)
        return source.get("source_name"), normalize_import_items(source.get("items", [])), source
    if payload.xlsx_base64:
        name, items = parse_xlsx(payload.xlsx_base64)
        return name, normalize_import_items(items), {"source_type": "xlsx", "source_url": None, "source_name": name}
    if payload.csv_text:
        name, items = parse_csv_items(payload.csv_text)
        return name, normalize_import_items(items), {"source_type": "csv", "source_url": None, "source_name": name}
    name, items = parse_json(payload.json_data or {})
    return name, normalize_import_items(items), {"source_type": "json", "source_url": None, "source_name": name}

