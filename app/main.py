import asyncio
import base64
import gzip
import ipaddress
import json
import math
import re
import socket
import sqlite3
import time
import uuid
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import httpx
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from .clients import (
    AIRecognitionClient,
    EmbyClient,
    MoviePilotClient,
    MTeamClient,
    NexusPHPClient,
    RSSClient,
    TMDBClient,
    TorznabClient,
    TransmissionClient,
)
from .candidate_policy import analyze as analyze_policy_candidate
from .candidate_policy import merge_custom_rules, normalized_policy, release_group_catalog
from .config import access_token, access_token_required, load_runtime_settings, save_runtime_settings, settings
from .cookiecloud import cookie_for_host, cookie_groups, decrypt_cookiecloud
from .database import config_values, connect, initialize, json_value, save_config, cleanup_old_data
from .list_sources import PlaylistSourceFetcher, parse_csv_items
from .security import extract_access_token, safe_error, sanitize_sensitive_text, token_matches


app = FastAPI(title="AutoList", version="0.75")
app.mount("/assets", StaticFiles(directory=Path(__file__).parent / "static"), name="assets")

# 公开路径在启用 AUTOLIST_ACCESS_TOKEN 时仍可访问；CookieCloud 协议路径另做 KEY 绑定。
AUTH_EXEMPT_PATHS = {"/", "/favicon.ico", "/api/health", "/cookiecloud", "/cookiecloud/"}
MAX_RUNNING_SEARCH_TASKS = 3
COOKIECLOUD_RATE_LIMIT = 10
COOKIECLOUD_RATE_WINDOW_SECONDS = 60
_cookiecloud_upload_times: list[float] = []


@app.get("/favicon.ico", include_in_schema=False)
async def favicon() -> FileResponse:
    return FileResponse(Path(__file__).parent / "static" / "favicon.svg", media_type="image/svg+xml")

# 下载 URL、Cookie 等短命敏感字段只保存在进程内，容器重启后会自然失效。
raw_candidates: dict[str, dict[str, Any]] = {}
running_tasks: dict[int, asyncio.Task[None]] = {}
running_recognition_tasks: dict[int, asyncio.Task[None]] = {}
running_library_tasks: dict[int, asyncio.Task[None]] = {}
running_automation_tasks: dict[int, asyncio.Task[None]] = {}
scheduler_task: asyncio.Task[None] | None = None
moviepilot_site_ids: dict[int, int] = {}
site_icon_cache: dict[int, tuple[bytes, str]] = {}
poster_cache: dict[str, tuple[bytes, str]] = {}


def enforce_cookiecloud_rate_limit() -> None:
    """Bound anonymous CookieCloud uploads to reduce disk-fill abuse."""
    now = time.time()
    cutoff = now - COOKIECLOUD_RATE_WINDOW_SECONDS
    while _cookiecloud_upload_times and _cookiecloud_upload_times[0] < cutoff:
        _cookiecloud_upload_times.pop(0)
    if len(_cookiecloud_upload_times) >= COOKIECLOUD_RATE_LIMIT:
        raise HTTPException(429, "CookieCloud 上传过于频繁，请稍后再试")
    _cookiecloud_upload_times.append(now)


def require_configured_cookiecloud_uuid(uuid_value: str) -> None:
    """Only the KEY configured in settings may read or write CookieCloud blobs."""
    configured = (settings.cookiecloud_key or "").strip()
    if not configured:
        raise HTTPException(503, "请先在设置中配置 CookieCloud 用户 KEY")
    if not token_matches(uuid_value, configured):
        raise HTTPException(403, "CookieCloud 用户 KEY 与服务端配置不匹配")


def enforce_search_task_capacity() -> None:
    active = sum(1 for task in running_tasks.values() if task and not task.done())
    if active >= MAX_RUNNING_SEARCH_TASKS:
        raise HTTPException(429, f"已有 {active} 个搜索任务在运行，请等待完成后再试")


@app.middleware("http")
async def access_token_middleware(request: Request, call_next):  # type: ignore[no-untyped-def]
    if not access_token_required():
        return await call_next(request)
    path = request.url.path
    if path in AUTH_EXEMPT_PATHS or path.startswith("/assets/") or path.startswith("/cookiecloud/"):
        return await call_next(request)
    provided = extract_access_token(request.headers.get("authorization"), request.headers.get("x-autolist-token"))
    if not token_matches(provided, access_token()):
        return JSONResponse({"detail": "需要有效的访问令牌"}, status_code=401)
    return await call_next(request)


class ImportPayload(BaseModel):
    name: str | None = Field(default=None, max_length=120)
    json_data: dict[str, Any] | list[dict[str, Any]] | None = None
    xlsx_base64: str | None = Field(default=None, max_length=50_000_000)
    csv_text: str | None = Field(default=None, max_length=20_000_000)
    source_url: str | None = Field(default=None, max_length=2048)
    limit: int = Field(default=5000, ge=1, le=10000)


class TaskPayload(BaseModel):
    playlist_id: int
    scope: str = Field(default="range", pattern="^(range|pending)$")
    count: int = Field(default=50, ge=1, le=10000)
    range_start: int = Field(default=1, ge=1)
    range_end: int = Field(default=1, ge=1)


class ConfigPayload(BaseModel):
    preferred_resolutions: str | None = None
    minimum_resolution: str | None = None
    preferred_codecs: str | None = None
    priority_groups: str | None = None
    secondary_groups: str | None = None
    fallback_groups: str | None = None
    allow_unknown_groups: bool | None = None
    candidate_limit: int | None = Field(default=None, ge=1, le=20)
    scoring_policy: dict[str, Any] | None = None
    candidate_policy: dict[str, Any] | None = None


class ScorePreviewPayload(BaseModel):
    title: str = Field(min_length=1, max_length=500)
    seeders: int = Field(default=0, ge=0)
    volume_factor: float = Field(default=1, ge=0)
    scoring_policy: dict[str, Any] | None = None
    candidate_policy: dict[str, Any] | None = None


class RuntimeSettingsPayload(BaseModel):
    mp_base_url: str = ""
    mp_api_key: str | None = None
    mp_timeout_seconds: float = Field(default=30, ge=3, le=300)
    emby_base_url: str = ""
    emby_api_key: str | None = None
    tmdb_api_key: str | None = None
    tmdb_language: str = "zh-CN"
    mdblist_api_key: str | None = None
    cookiecloud_key: str = ""
    cookiecloud_password: str | None = None
    cookiecloud_forward_moviepilot: bool = True
    outbound_proxy_url: str | None = None
    tmdb_proxy_enabled: bool = False
    pt_proxy_enabled: bool = False
    ai_base_url: str = ""
    ai_api_key: str | None = None
    ai_model: str = ""
    tr_base_url: str = ""
    tr_username: str = ""
    tr_password: str | None = None
    dashboard_random_posters: bool = False


class SitePayload(BaseModel):
    name: str = Field(min_length=1, max_length=80)
    base_url: str = ""
    api_key: str | None = None
    cookie: str | None = None
    user_agent: str = ""
    priority: int = Field(default=100, ge=1, le=999)
    timeout_seconds: int = Field(default=30, ge=3, le=300)
    rss_url: str = ""
    icon_url: str = ""
    proxy: bool = False
    render: bool = False
    limit_interval: int | None = Field(default=None, ge=1)
    limit_count: int | None = Field(default=None, ge=1)
    enabled: bool = True
    search_enabled: bool = False


class SiteCookiePayload(BaseModel):
    username: str = ""
    password: str = ""
    code: str = ""


class CookieCloudUploadPayload(BaseModel):
    uuid: str = Field(min_length=5, max_length=128, pattern=r"^[A-Za-z0-9_-]+$")
    encrypted: str = Field(min_length=16, max_length=32_000_000)
    crypto_type: str = Field(default="legacy", pattern=r"^(legacy|aes-128-cbc-fixed)$")


class PlaylistUpdatePayload(BaseModel):
    name: str = Field(min_length=1, max_length=120)


class PlaylistOrderPayload(BaseModel):
    ids: list[int]


class PlaylistAutomationPayload(BaseModel):
    enabled: bool = False
    auto_cart: bool = False
    batch_size: int = Field(default=50, ge=1, le=200)


class PlaylistSyncPayload(BaseModel):
    enabled: bool = False
    interval_hours: int = Field(default=24, ge=1, le=720)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def cookiecloud_file(uuid_value: str) -> Path:
    if not re.fullmatch(r"[A-Za-z0-9_-]{5,128}", uuid_value):
        raise HTTPException(422, "CookieCloud 用户 KEY 格式无效")
    directory = Path(settings.data_dir) / "cookiecloud"
    directory.mkdir(parents=True, exist_ok=True)
    return directory / f"{uuid_value}.json"


def stored_cookiecloud_payload() -> dict[str, Any]:
    if not settings.cookiecloud_key or not settings.cookiecloud_password:
        raise HTTPException(422, "请先在设置中配置 CookieCloud 用户 KEY 与端对端密码")
    path = cookiecloud_file(settings.cookiecloud_key)
    if not path.exists():
        raise HTTPException(404, "尚未收到 Chrome CookieCloud 数据，请在扩展中执行一次同步")
    try:
        stored = json.loads(path.read_text(encoding="utf-8"))
        return decrypt_cookiecloud(settings.cookiecloud_key, settings.cookiecloud_password, stored["encrypted"], stored.get("crypto_type", "legacy"))
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, f"CookieCloud 解密失败：{safe_error(exc)}") from exc


def rows_to_dicts(rows: list[sqlite3.Row]) -> list[dict[str, Any]]:
    return [dict(row) for row in rows]


def secret_free(value: Any) -> Any:
    if isinstance(value, list):
        return [secret_free(item) for item in value]
    if not isinstance(value, dict):
        return value
    blocked = re.compile(r"(url|cookie|passkey|token|authorization|api.?key|header)", re.I)
    return {key: secret_free(item) for key, item in value.items() if not blocked.search(key)}


def raster_image_media_type(content: bytes) -> str | None:
    """Identify supported raster icon formats by magic bytes; remote SVG/HTML is never re-served."""
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if content.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if content.startswith((b"GIF87a", b"GIF89a")):
        return "image/gif"
    if content.startswith(b"RIFF") and content[8:12] == b"WEBP":
        return "image/webp"
    if content.startswith(b"\x00\x00\x01\x00"):
        return "image/x-icon"
    return None


MAX_COOKIECLOUD_BODY = 40 * 1024 * 1024


def decode_cookiecloud_body(content: bytes, content_encoding: str, limit: int = MAX_COOKIECLOUD_BODY) -> bytes:
    """Bound compressed and expanded CookieCloud uploads before JSON validation."""
    if len(content) > limit:
        raise HTTPException(413, "CookieCloud 上传数据过大")
    if "gzip" in content_encoding.lower():
        try:
            with gzip.GzipFile(fileobj=BytesIO(content)) as compressed:
                content = compressed.read(limit + 1)
        except (OSError, EOFError) as exc:
            raise HTTPException(422, "CookieCloud gzip 数据无效") from exc
    if len(content) > limit:
        raise HTTPException(413, "CookieCloud 上传数据过大")
    return content


async def validate_remote_icon_url(source: str, site_base_url: str) -> None:
    """Allow configured-site icons while preventing cross-host requests into private networks."""
    parsed = urlparse(source)
    base = urlparse(site_base_url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise RuntimeError("图标地址无效")
    if parsed.hostname.lower() == (base.hostname or "").lower():
        return
    try:
        addresses = await asyncio.to_thread(socket.getaddrinfo, parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80))
    except socket.gaierror as exc:
        raise RuntimeError("图标域名无法解析") from exc
    if not addresses or any(not ipaddress.ip_address(address[4][0]).is_global for address in addresses):
        raise RuntimeError("图标地址不允许访问内网")


def first_value(data: dict[str, Any], names: tuple[str, ...], default: Any = None) -> Any:
    for name in names:
        if data.get(name) not in (None, ""):
            return data[name]
    return default


def resource_fingerprint(title: str, size: int | None = None) -> str:
    normalized = re.sub(r"\b(?:free|2x|50%|30%)\b", "", title.lower())
    normalized = re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", normalized)
    size_bucket = round(int(size or 0) / (256 * 1024 * 1024)) if size else 0
    return f"{normalized}:{size_bucket}"


def volume_factor_value(value: Any) -> float:
    if value is None:
        return 1.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().lower()
    if text in ("free", "免费", "freeleech"):
        return 0.0
    if text.endswith("%"):
        try:
            return float(text[:-1]) / 100
        except ValueError:
            return 1.0
    try:
        return float(text)
    except ValueError:
        return 1.0


def parse_xlsx(encoded: str) -> tuple[str | None, list[dict[str, Any]]]:
    workbook = None
    try:
        content = base64.b64decode(encoded, validate=True)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(422, f"无法读取 xlsx：{safe_error(exc)}") from exc
    try:
        sheet = workbook.active
        first_row = next(sheet.iter_rows(min_row=1, max_row=1), None)
        if not first_row:
            raise HTTPException(422, "xlsx 为空，请确认文件包含表头和影片数据")
        headers = [str(cell.value or "").strip() for cell in first_row]
        required = {"总排名", "IMDb ID", "英文/原片名", "年份", "中文译名"}
        if not required.issubset(headers):
            raise HTTPException(422, "xlsx 缺少必需列：总排名、IMDb ID、英文/原片名、年份、中文译名")
        positions = {header: index for index, header in enumerate(headers)}
        items: list[dict[str, Any]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            title = row[positions["英文/原片名"]]
            if not title:
                continue
            year = row[positions["年份"]]
            items.append({
                "rank_no": row[positions["总排名"]],
                "imdb_id": row[positions["IMDb ID"]],
                "original_title": str(title).strip(),
                "year": int(year) if str(year or "").isdigit() else None,
                "chinese_title": row[positions["中文译名"]],
            })
        return sheet.title, items
    finally:
        workbook.close()


def parse_json(data: dict[str, Any] | list[dict[str, Any]]) -> tuple[str | None, list[dict[str, Any]]]:
    source = data if isinstance(data, dict) else {"films": data}
    films = source.get("films") or source.get("items") or []
    if not isinstance(films, list):
        raise HTTPException(422, "JSON 必须包含 films 或 items 数组")
    items: list[dict[str, Any]] = []
    for index, film in enumerate(films, start=1):
        if not isinstance(film, dict):
            continue
        title = first_value(film, ("original_title", "title", "english_title", "英文/原片名"))
        if not title:
            continue
        year = first_value(film, ("year", "年份"))
        items.append({
            "rank_no": first_value(film, ("rank_no", "rank", "总排名"), index),
            "imdb_id": first_value(film, ("imdb_id", "imdbId", "imdb", "IMDb ID")),
            "original_title": str(title).strip(),
            "year": int(year) if str(year or "").isdigit() else None,
            "chinese_title": first_value(film, ("chinese_title", "cn_title", "中文译名")),
            "tmdb_id": first_value(film, ("tmdb_id", "tmdbId", "tmdb")),
        })
    return source.get("listName") or source.get("name"), items


def normalize_import_items(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in items:
        title = str(item.get("original_title") or "").strip()
        imdb_id = str(item.get("imdb_id") or "").strip() or None
        tmdb_raw = item.get("tmdb_id")
        tmdb_id = int(tmdb_raw) if str(tmdb_raw or "").isdigit() else None
        year_raw = item.get("year")
        year = int(year_raw) if str(year_raw or "").isdigit() else None
        if not title:
            continue
        normalized_title = re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", title.lower())
        key = f"tmdb:{tmdb_id}" if tmdb_id else (f"imdb:{imdb_id.lower()}" if imdb_id else f"title:{normalized_title}:{year or ''}")
        if key in seen:
            continue
        seen.add(key)
        result.append({"rank_no": len(result) + 1, "imdb_id": imdb_id, "original_title": title, "year": year,
                       "chinese_title": item.get("chinese_title"), "tmdb_id": tmdb_id})
    return result


async def resolve_import(payload: ImportPayload) -> tuple[str | None, list[dict[str, Any]], dict[str, Any]]:
    if payload.source_url:
        source = await PlaylistSourceFetcher().fetch(payload.source_url, payload.limit)
        return source.get("source_name"), normalize_import_items(source.get("items", [])), source
    if payload.xlsx_base64:
        name, items = parse_xlsx(payload.xlsx_base64)
        return name, normalize_import_items(items), {"source_type": "xlsx", "source_url": None, "source_name": name}
    if payload.csv_text:
        name, items = parse_csv_items(payload.csv_text)
        return name, normalize_import_items(items), {"source_type": "csv", "source_url": None, "source_name": name}
    name, items = parse_json(payload.json_data or {})
    return name, normalize_import_items(items), {"source_type": "json", "source_url": None, "source_name": name}




def analyze_candidate(title: str, index: int, config: dict[str, str], torrent: dict[str, Any] | None = None) -> dict[str, Any]:
    try:
        policy = json.loads(config.get("candidate_policy") or "{}")
    except (TypeError, json.JSONDecodeError):
        policy = {}
    return analyze_policy_candidate(title, index, policy, torrent)


def extract_contexts(response: Any) -> list[dict[str, Any]]:
    if isinstance(response, dict):
        data = response.get("data", response.get("result", []))
    else:
        data = response
    if isinstance(data, dict):
        data = data.get("contexts", data.get("items", []))
    return [item for item in data if isinstance(item, dict)] if isinstance(data, list) else []


def extract_pair(context: dict[str, Any]) -> tuple[dict[str, Any] | None, dict[str, Any]] | None:
    media = context.get("media_info") or context.get("media")
    torrent = context.get("torrent_info") or context.get("torrent")
    if isinstance(torrent, dict):
        return media if isinstance(media, dict) else None, torrent
    if any(context.get(key) for key in ("enclosure", "download_url", "magnet")):
        return None, context
    return None


def select_tmdb_match(options: list[dict[str, Any]], title: str, year: int | None) -> dict[str, Any] | None:
    if not options:
        return None
    normalized = re.sub(r"[^a-z0-9]+", " ", title.lower()).strip()
    for item in options:
        item_year = str(item.get("release_date") or "")[:4]
        names = (item.get("title"), item.get("original_title"))
        normalized_names = {re.sub(r"[^a-z0-9]+", " ", str(name).lower()).strip() for name in names if name}
        if normalized in normalized_names and (not year or item_year == str(year)):
            return item
    same_year = [item for item in options if not year or str(item.get("release_date") or "")[:4] == str(year)]
    return same_year[0] if same_year else options[0]


async def recognize_movie(title: str, year: int | None, imdb_id: str | None = None) -> dict[str, Any] | None:
    tmdb = TMDBClient()
    async def enriched(match: dict[str, Any] | None) -> dict[str, Any] | None:
        if not match:
            return None
        result = dict(match)
        resolved_imdb = imdb_id
        if not resolved_imdb and result.get("id"):
            try:
                resolved_imdb = (await tmdb.movie_external_ids(int(result["id"]))).get("imdb_id")
            except Exception:
                resolved_imdb = None
        result["imdb_id"] = resolved_imdb
        return result
    if imdb_id:
        match = select_tmdb_match(await tmdb.find_by_imdb(imdb_id), title, year)
        if match:
            return await enriched(match)
    options = await tmdb.search_movie(title, year)
    match = select_tmdb_match(options, title, year)
    if match:
        return await enriched(match)
    suggestion = await AIRecognitionClient().suggest(title, year)
    if not suggestion:
        return None
    options = await tmdb.search_movie(str(suggestion.get("original_title") or suggestion.get("title") or title), suggestion.get("year") or year)
    return await enriched(select_tmdb_match(
        options, str(suggestion.get("original_title") or suggestion.get("title") or title), suggestion.get("year") or year,
    ))


def tmdb_item_values(media: dict[str, Any], fallback_imdb: str | None = None) -> tuple[Any, ...]:
    release_year = str(media.get("release_date") or "")[:4]
    return (
        int(media["id"]), str(media.get("title") or "").strip() or None,
        str(media.get("original_title") or "").strip() or None,
        int(release_year) if release_year.isdigit() else None,
        media.get("imdb_id") or fallback_imdb, utc_now(),
    )


def persist_tmdb_item(item_id: int, media: dict[str, Any], fallback_imdb: str | None = None) -> None:
    values = tmdb_item_values(media, fallback_imdb)
    with connect() as conn:
        conn.execute(
            """UPDATE playlist_items
               SET tmdb_id=?,tmdb_title=?,tmdb_original_title=?,tmdb_year=?,tmdb_imdb_id=?,tmdb_checked_at=?
               WHERE id=?""",
            (*values, item_id),
        )


def canonical_item_title(item: sqlite3.Row | dict[str, Any]) -> str:
    return str(item["tmdb_title"] or item["chinese_title"] or item["original_title"])


def canonical_item_original_title(item: sqlite3.Row | dict[str, Any]) -> str:
    return str(item["tmdb_original_title"] or item["original_title"])


def canonical_item_year(item: sqlite3.Row | dict[str, Any]) -> int | None:
    value = item["tmdb_year"] or item["year"]
    return int(value) if value else None


def normalized_title_text(value: Any) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", str(value or "").casefold())


def title_tokens(value: Any) -> set[str]:
    return {
        token for token in re.sub(r"[^a-z0-9\u4e00-\u9fff]+", " ", str(value or "").casefold()).split()
        if len(token) > 1 or token.isdigit()
    }


TITLE_STOP_WORDS = {
    "a", "an", "and", "at", "by", "da", "das", "de", "del", "der", "die", "di", "dos",
    "for", "from", "in", "la", "le", "les", "of", "on", "or", "the", "to", "un", "una",
    "upon", "with", "once", "time",
}


def informative_title_tokens(value: Any) -> set[str]:
    return title_tokens(value) - TITLE_STOP_WORDS


def strict_torrent_matches_item(item: sqlite3.Row | dict[str, Any], torrent_title: str) -> bool:
    """Require a meaningful title token so shared words cannot identify another movie."""
    candidate = normalized_title_text(torrent_title)
    candidate_tokens = title_tokens(torrent_title)
    if not candidate or not candidate_tokens:
        return False
    variants = [
        item["tmdb_original_title"] if item["tmdb_original_title"] else None,
        item["tmdb_title"] if item["tmdb_title"] else None,
        item["original_title"], item["chinese_title"] if item["chinese_title"] else None,
    ]
    for variant in variants:
        normalized = normalized_title_text(variant)
        tokens = informative_title_tokens(variant)
        if tokens and tokens.issubset(candidate_tokens):
            return True
        if normalized and normalized in candidate and (not tokens or not any(token.isdigit() for token in tokens)):
            return True
    return False


def torrent_matches_item(item: sqlite3.Row | dict[str, Any], torrent_title: str) -> bool:
    candidate = normalized_title_text(torrent_title)
    if not candidate:
        return False
    variants = [
        item["tmdb_original_title"] if item["tmdb_original_title"] else None,
        item["tmdb_title"] if item["tmdb_title"] else None,
        item["original_title"], item["chinese_title"] if item["chinese_title"] else None,
    ]
    candidate_tokens = title_tokens(torrent_title)
    for variant in variants:
        normalized = normalized_title_text(variant)
        tokens = title_tokens(variant)
        if tokens and tokens.issubset(candidate_tokens):
            return True
        if normalized and normalized in candidate and (not tokens or not any(token.isdigit() for token in tokens)):
            return True
    return False


def candidate_identity(item: sqlite3.Row | dict[str, Any], media: dict[str, Any], torrent_title: str) -> tuple[bool, str | None]:
    target_year = str(media.get("year") or canonical_item_year(item) or "").strip()
    years = set(re.findall(r"(?<!\d)(?:19|20)\d{2}(?!\d)", torrent_title))
    if target_year and years and target_year not in years:
        return False, f"年份不匹配：目标 {target_year}，资源包含 {', '.join(sorted(years))}"
    if re.search(r"(?i)(?:trilogy|collection|box[ ._-]*set|complete|pack|合集|系列|全集)", torrent_title):
        return False, "疑似合集或系列资源"
    if not strict_torrent_matches_item(item, torrent_title):
        return False, "片名不匹配：资源片名与目标影片不一致"
    return True, None


def normalized_download_name(value: Any) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", str(value or "").casefold())


def is_transmission_downloading(torrent: dict[str, Any]) -> bool:
    status = torrent.get("status")
    try:
        status_value = int(status)
    except (TypeError, ValueError):
        status_value = -1
    try:
        percent_done = float(torrent.get("percentDone") or 0)
    except (TypeError, ValueError):
        percent_done = 0
    return status_value in {1, 2, 3, 4} and percent_done < 1


DOWNLOAD_LIFECYCLE_LABELS = {
    "submitted": "已提交",
    "downloading": "下载中",
    "pending_confirmation": "待确认",
    "pending_library": "待入库",
    "organized": "已整理/已入库",
    "failed": "失败",
}


def _history_item_snapshot(row: dict[str, Any]) -> dict[str, Any] | None:
    item_id = row.get("resolved_playlist_item_id")
    if item_id is None:
        return None
    return {
        "id": item_id,
        "original_title": row.get("playlist_original_title"),
        "chinese_title": row.get("playlist_chinese_title"),
        "year": row.get("playlist_year"),
        "tmdb_title": row.get("playlist_tmdb_title"),
        "tmdb_original_title": row.get("playlist_tmdb_original_title"),
        "tmdb_year": row.get("playlist_tmdb_year"),
    }


def _history_torrent_matches(row: dict[str, Any], torrent: dict[str, Any]) -> bool:
    torrent_hash = str(torrent.get("hashString") or "").strip().casefold()
    submission_hash = str(row.get("submission_hash") or "").strip().casefold()
    if torrent_hash and submission_hash and torrent_hash == submission_hash:
        return True
    torrent_name = normalized_download_name(first_value(torrent, ("name", "torrent_name"), ""))
    history_name = normalized_download_name(row.get("torrent_name"))
    if torrent_name and history_name and torrent_name == history_name:
        return True
    item = row.get("_playlist_item")
    return bool(item and strict_torrent_matches_item(item, str(first_value(torrent, ("name", "torrent_name"), ""))))


def _active_history_matches(
    histories: list[dict[str, Any]], torrents: list[dict[str, Any]],
) -> tuple[set[int], set[int]]:
    matched: set[int] = set()
    ambiguous: set[int] = set()
    eligible = [row for row in histories if bool(row.get("success"))]
    active_torrents = [torrent for torrent in torrents if is_transmission_downloading(torrent)]
    for torrent in active_torrents:
        torrent_hash = str(torrent.get("hashString") or "").strip().casefold()
        hash_matches = [
            row for row in eligible
            if torrent_hash and str(row.get("submission_hash") or "").strip().casefold() == torrent_hash
        ]
        if len(hash_matches) == 1:
            matched.add(int(hash_matches[0]["id"]))
            continue
        if len(hash_matches) > 1:
            ambiguous.update(int(row["id"]) for row in hash_matches)
            continue
        name = normalized_download_name(first_value(torrent, ("name", "torrent_name"), ""))
        name_matches = [
            row for row in eligible
            if name and normalized_download_name(row.get("torrent_name")) == name
        ]
        if len(name_matches) == 1:
            matched.add(int(name_matches[0]["id"]))
            continue
        if len(name_matches) > 1:
            ambiguous.update(int(row["id"]) for row in name_matches)
            continue
        identity_matches = [row for row in eligible if _history_torrent_matches(row, torrent)]
        if len(identity_matches) == 1:
            matched.add(int(identity_matches[0]["id"]))
        elif len(identity_matches) > 1:
            ambiguous.update(int(row["id"]) for row in identity_matches)
    return matched, ambiguous


def _project_history_state(
    row: dict[str, Any], matched_ids: set[int], ambiguous_ids: set[int],
    transmission_error: bool, checked_at: str,
) -> dict[str, Any]:
    history_id = int(row["id"])
    library_state = str(row.get("playlist_library_state") or "unknown")
    if not bool(row.get("success")):
        lifecycle_status, source, reason, next_action = (
            "failed", "AutoList/MoviePilot", row.get("message") or "提交失败", "查看失败原因",
        )
        status_checked_at = row.get("created_at")
    elif library_state == "in_library":
        lifecycle_status, source, reason, next_action = (
            "organized", "Emby", "Emby 已找到实体媒体", "无需操作",
        )
        status_checked_at = row.get("playlist_library_checked_at") or checked_at
    elif library_state == "strm":
        lifecycle_status, source, reason, next_action = (
            "pending_library", "Emby", "Emby 已找到 .strm，实体媒体尚未确认", "刷新 Emby 状态",
        )
        status_checked_at = row.get("playlist_library_checked_at") or checked_at
    elif history_id in matched_ids:
        lifecycle_status, source, reason, next_action = (
            "downloading", "Transmission", "Transmission 正在下载", "等待下游确认",
        )
        status_checked_at = checked_at
    elif transmission_error or history_id in ambiguous_ids or (
        library_state == "not_found" and row.get("playlist_library_checked_at")
    ):
        lifecycle_status, source, reason, next_action = (
            "pending_confirmation", "Transmission/Emby", "暂时无法确认下游状态", "稍后刷新状态",
        )
        status_checked_at = row.get("playlist_library_checked_at") or checked_at
    else:
        lifecycle_status, source, reason, next_action = (
            "submitted", "MoviePilot", "已提交给 MoviePilot，等待下游服务确认", "等待下游确认",
        )
        status_checked_at = row.get("created_at")
    item = {key: value for key, value in row.items() if not key.startswith("_")}
    for key in (
        "resolved_playlist_item_id", "playlist_original_title", "playlist_chinese_title", "playlist_year",
        "playlist_tmdb_title", "playlist_tmdb_original_title", "playlist_tmdb_year", "playlist_library_state",
        "playlist_library_checked_at", "playlist_item_id", "submission_hash",
    ):
        item.pop(key, None)
    item["lifecycle_status"] = lifecycle_status
    item["status_label"] = DOWNLOAD_LIFECYCLE_LABELS[lifecycle_status]
    item["status_source"] = source
    item["status_reason"] = sanitize_sensitive_text(str(reason), 500)
    item["status_checked_at"] = status_checked_at
    item["next_action"] = next_action
    if item.get("message"):
        item["message"] = sanitize_sensitive_text(item["message"])
    return item


async def projected_download_history(limit: int = 200) -> list[dict[str, Any]]:
    safe_limit = max(1, min(limit, 200))
    with connect() as conn:
        rows = conn.execute(
            """SELECT h.*, COALESCE(h.playlist_item_id,c.playlist_item_id) AS resolved_playlist_item_id,
                      p.original_title AS playlist_original_title,p.chinese_title AS playlist_chinese_title,
                      p.year AS playlist_year,p.tmdb_title AS playlist_tmdb_title,
                      p.tmdb_original_title AS playlist_tmdb_original_title,p.tmdb_year AS playlist_tmdb_year,
                      p.library_state AS playlist_library_state,p.library_checked_at AS playlist_library_checked_at
               FROM download_history h
               LEFT JOIN candidates c ON c.id=h.candidate_id
               LEFT JOIN playlist_items p ON p.id=COALESCE(h.playlist_item_id,c.playlist_item_id)
               ORDER BY h.id DESC LIMIT ?""",
            (safe_limit,),
        ).fetchall()
    histories = rows_to_dicts(rows)
    if not histories:
        return []
    for row in histories:
        row["_playlist_item"] = _history_item_snapshot(row)
    transmission_error = False
    try:
        torrents = await asyncio.wait_for(TransmissionClient().current_downloads(), timeout=6)
    except Exception:
        torrents = []
        transmission_error = True
    matched_ids, ambiguous_ids = _active_history_matches(histories, torrents)
    checked_at = utc_now()
    return [_project_history_state(row, matched_ids, ambiguous_ids, transmission_error, checked_at) for row in histories]


async def searchable_playlist_items(playlist_id: int, limit: int | None = None) -> dict[str, Any]:
    with connect() as conn:
        playlist = conn.execute("SELECT id,name FROM playlists WHERE id=?", (playlist_id,)).fetchone()
        if not playlist:
            raise HTTPException(404, "片单不存在")
        stats = conn.execute(
            """SELECT COUNT(*) AS total,
                      SUM(CASE WHEN library_state='in_library' THEN 1 ELSE 0 END) AS in_library
               FROM playlist_items WHERE playlist_id=?""", (playlist_id,),
        ).fetchone()
        rows = list(conn.execute(
            "SELECT * FROM playlist_items WHERE playlist_id=? AND library_state!='in_library' ORDER BY rank_no",
            (playlist_id,),
        ).fetchall())
        history_rows = conn.execute(
            """SELECT h.torrent_name,h.title,c.title AS candidate_title,c.playlist_item_id
               FROM download_history h
               LEFT JOIN candidates c ON c.id=h.candidate_id
               JOIN playlist_items p ON p.id=c.playlist_item_id
               WHERE p.playlist_id=?""", (playlist_id,),
        ).fetchall()
        candidate_rows = conn.execute(
            """SELECT NULL AS torrent_name,c.title,c.title AS candidate_title,c.playlist_item_id FROM candidates c
               JOIN playlist_items p ON p.id=c.playlist_item_id
               WHERE p.playlist_id=?""", (playlist_id,),
        ).fetchall()
    try:
        torrents = await asyncio.wait_for(TransmissionClient().current_downloads(), timeout=6)
    except Exception:
        torrents = []
    downloading = [torrent for torrent in torrents if is_transmission_downloading(torrent)]
    download_names = {normalized_download_name(first_value(torrent, ("name", "torrent_name"), "")) for torrent in downloading}
    related_names: dict[str, int] = {}
    for row in list(history_rows) + list(candidate_rows):
        item_id = row["playlist_item_id"]
        for name in (row["torrent_name"], row["title"], row["candidate_title"]):
            normalized = normalized_download_name(name)
            if normalized:
                related_names[normalized] = int(item_id)
    downloading_ids: set[int] = set()
    for name in download_names:
        if name in related_names:
            downloading_ids.add(related_names[name])
    for torrent in downloading:
        torrent_title = str(first_value(torrent, ("name", "torrent_name"), ""))
        for item in rows:
            if int(item["id"]) not in downloading_ids and torrent_matches_item(item, torrent_title):
                downloading_ids.add(int(item["id"]))
    queue = [item for item in rows if int(item["id"]) not in downloading_ids]
    selected = queue[:limit] if limit is not None else queue
    return {
        "playlist_id": playlist_id, "playlist_name": playlist["name"],
        "total_count": int(stats["total"] or 0), "in_library_count": int(stats["in_library"] or 0),
        "downloading_count": len(downloading),
        "pending_count": len(queue), "items": rows_to_dicts(selected),
    }


def build_search_queries(item: sqlite3.Row, media: dict[str, Any]) -> list[tuple[str, str | None, str]]:
    """Return a bounded IMDb/title search plan, preserving TMDB as the authority."""
    item_keys = item.keys() if hasattr(item, "keys") else ()
    tmdb_imdb_id = item["tmdb_imdb_id"] if "tmdb_imdb_id" in item_keys else None
    item_imdb_id = item["imdb_id"] if "imdb_id" in item_keys else None
    imdb_id = str(media.get("imdb_id") or tmdb_imdb_id or item_imdb_id or "").strip() or None
    year = str(media.get("year") or canonical_item_year(item) or "").strip()
    titles = [
        str(media.get("original_title") or "").strip(),
        str(media.get("title") or "").strip(),
        str(item["original_title"] or "").strip(),
        str(item["chinese_title"] or "").strip(),
    ]
    queries: list[tuple[str, str | None, str]] = []
    if imdb_id:
        queries.append((titles[0] or titles[1], imdb_id, f"IMDb {imdb_id}"))
    seen: set[str] = set()
    for title in titles:
        key = re.sub(r"\W+", "", title).casefold()
        if not title or key in seen:
            continue
        seen.add(key)
        keyword = title if year and re.search(rf"(?:^|\D){re.escape(year)}(?:\D|$)", title) else f"{title} {year}".strip()
        queries.append((keyword, None, keyword))
    return queries[:4]


def domain_match(domain: str, base_url: str) -> bool:
    host = (urlparse(base_url).hostname or "").lower().removeprefix("www.")
    value = domain.lower().removeprefix("www.")
    return host == value or host.endswith(f".{value}") or value.endswith(f".{host}")


def resolve_site_adapter(base_url: str, rss_url: str = "") -> str:
    """Keep adapter details out of the UI; select known special protocols server-side."""
    if rss_url.strip():
        return "rss"
    normalized = base_url.lower()
    if "m-team" in normalized or "mteam" in normalized:
        return "mteam"
    if "torznab" in normalized or "api?t=" in normalized or "t=caps" in normalized:
        return "torznab"
    return "nexusphp"


async def moviepilot_site_snapshot() -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    client = MoviePilotClient()
    if not settings.mp_base_url or not settings.mp_api_key:
        return [], [], []
    sites, statistics, users = await asyncio.gather(client.sites(), client.site_statistics(), client.site_user_data())
    return (sites if isinstance(sites, list) else [], statistics if isinstance(statistics, list) else [], users if isinstance(users, list) else [])


async def test_site_config(site: dict[str, Any]) -> dict[str, Any]:
    try:
        if site["adapter"] == "mteam":
            result = await MTeamClient().check(site)
        elif site["adapter"] == "nexusphp":
            result = await NexusPHPClient().check(site)
        elif site["adapter"] == "rss":
            result = await RSSClient().check(site)
        else:
            torrents = await TorznabClient().search(site, "AutoListConnectionProbe")
            result = {"ok": True, "message": f"Torznab 可用，探测返回 {len(torrents)} 条"}
        status, message = "ok", sanitize_sensitive_text(result.get("message") or "连接正常")
    except Exception as exc:
        status, message = "error", safe_error(exc)
    with connect() as conn:
        conn.execute("UPDATE pt_sites SET last_status=?,last_message=?,last_tested_at=? WHERE id=?", (status, message[:500], utc_now(), site["id"]))
    return {"id": site["id"], "name": site["name"], "ok": status == "ok", "message": message}


async def library_details(
    emby: EmbyClient, title: str, year: int | None, tmdb_id: int | None = None, imdb_id: str | None = None,
) -> tuple[str, str | None, str | None]:
    try:
        state, item = await emby.library_match(title, year, tmdb_id, imdb_id)
        item_id = str(item.get("Id") or "") or None if item else None
        image_tag = str((item.get("ImageTags") or {}).get("Primary") or "") or None if item else None
        return state, item_id, image_tag
    except Exception:
        return "unknown", None, None


async def library_state(
    emby: EmbyClient, title: str, year: int | None, tmdb_id: int | None = None, imdb_id: str | None = None,
) -> str:
    state, _, _ = await library_details(emby, title, year, tmdb_id, imdb_id)
    return state


def update_library_task(task_id: int, **values: Any) -> None:
    if not set(values).issubset({"status", "completed", "in_library", "strm", "error_message"}):
        raise ValueError("无效的入库任务字段")
    values["updated_at"] = utc_now()
    assignments = ", ".join(f"{key}=?" for key in values)
    with connect() as conn:
        # Dynamic column names are restricted by the allowlist above.
        conn.execute(f"UPDATE library_scan_tasks SET {assignments} WHERE id=?", (*values.values(), task_id))  # nosec B608


async def run_library_scan(task_id: int) -> None:
    try:
        with connect() as conn:
            task = conn.execute("SELECT * FROM library_scan_tasks WHERE id=?", (task_id,)).fetchone()
            if not task:
                return
            items = rows_to_dicts(conn.execute("SELECT * FROM playlist_items WHERE playlist_id=? ORDER BY rank_no", (task["playlist_id"],)).fetchall())
        update_library_task(task_id, status="running")
        semaphore = asyncio.Semaphore(6)
        emby = EmbyClient()

        async def inspect(item: dict[str, Any]) -> tuple[int, str, str | None, str | None]:
            async with semaphore:
                state, emby_item_id, image_tag = await library_details(
                    emby, canonical_item_title(item), canonical_item_year(item), item["tmdb_id"],
                    item["tmdb_imdb_id"] or item["imdb_id"],
                )
                return int(item["id"]), state, emby_item_id, image_tag

        completed = in_library = strm = 0
        for future in asyncio.as_completed([inspect(item) for item in items]):
            item_id, state, emby_item_id, image_tag = await future
            completed += 1
            in_library += state == "in_library"
            strm += state == "strm"
            with connect() as conn:
                conn.execute(
                    "UPDATE playlist_items SET library_state=?,library_checked_at=?,emby_item_id=?,emby_image_tag=? WHERE id=?",
                    (state, utc_now(), emby_item_id, image_tag, item_id),
                )
            update_library_task(task_id, completed=completed, in_library=in_library, strm=strm)
        update_library_task(task_id, status="completed", completed=completed, in_library=in_library, strm=strm)
    except asyncio.CancelledError:
        update_library_task(task_id, status="cancelled")
        raise
    except Exception as exc:
        update_library_task(task_id, status="failed", error_message=safe_error(exc))
    finally:
        running_library_tasks.pop(task_id, None)


def update_task(task_id: int, **values: Any) -> None:
    if not set(values).issubset({"status", "completed", "matched", "error_message"}):
        raise ValueError("无效的搜索任务字段")
    values["updated_at"] = utc_now()
    assignments = ", ".join(f"{key}=?" for key in values)
    with connect() as conn:
        # Dynamic column names are restricted by the allowlist above.
        conn.execute(f"UPDATE search_tasks SET {assignments} WHERE id=?", (*values.values(), task_id))  # nosec B608


def task_log(task_id: int, level: str, stage: str, message: str) -> None:
    with connect() as conn:
        conn.execute(
            "INSERT INTO search_task_logs(task_id,level,stage,message,created_at) VALUES(?,?,?,?,?)",
            (task_id, level, stage, sanitize_sensitive_text(message, 1000), utc_now()),
        )


SEARCH_ERROR_MESSAGES = {
    "dns_error": "无法解析站点地址，请检查域名或 DNS 设置",
    "connect_error": "无法连接站点，请检查地址和网络",
    "timeout": "站点响应超时，请稍后重试",
    "auth_error": "站点认证失败，请检查登录信息",
    "rate_limit": "站点请求过于频繁，请稍后重试",
    "http_error": "站点返回异常，请稍后重试或检查站点状态",
    "parse_error": "站点返回内容无法解析，请重试或检查站点适配",
    "error": "站点搜索失败，请查看日志后重试",
}


def classify_search_error(exc: Exception) -> tuple[str, str]:
    technical = safe_error(exc).casefold()
    if isinstance(exc, httpx.TimeoutException) or "timed out" in technical or "timeout" in technical:
        return "timeout", SEARCH_ERROR_MESSAGES["timeout"]
    if isinstance(exc, httpx.HTTPStatusError):
        status = exc.response.status_code
        if status in {401, 403}:
            return "auth_error", SEARCH_ERROR_MESSAGES["auth_error"]
        if status == 429:
            return "rate_limit", SEARCH_ERROR_MESSAGES["rate_limit"]
        return "http_error", f"站点返回 HTTP {status}，请稍后重试"
    if isinstance(exc, httpx.NetworkError) or re.search(
        r"(?:name or service not known|nodename nor servname|temporary failure in name resolution|gaierror)",
        technical,
    ):
        if re.search(r"(?:name or service not known|nodename nor servname|name resolution|gaierror)", technical):
            return "dns_error", SEARCH_ERROR_MESSAGES["dns_error"]
        return "connect_error", SEARCH_ERROR_MESSAGES["connect_error"]
    if isinstance(exc, (ValueError, TypeError, json.JSONDecodeError)) or re.search(r"(?:parse|xml|json)", technical):
        return "parse_error", SEARCH_ERROR_MESSAGES["parse_error"]
    return "error", SEARCH_ERROR_MESSAGES["error"]


def record_search_attempt(
    task_id: int, item_id: int, site: dict[str, Any], attempt_no: int, status: str,
    result_count: int, duration_ms: int, error_code: str | None = None, error_message: str | None = None,
    query_count: int = 1,
) -> None:
    now = utc_now()
    with connect() as conn:
        conn.execute(
            """INSERT INTO search_attempts(
                 task_id,playlist_item_id,site_id,site_name,attempt_no,status,result_count,duration_ms,
                 error_code,error_message,query_count,created_at,finished_at
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (task_id, item_id, site["id"], site["name"], attempt_no, status, result_count,
             duration_ms, error_code, sanitize_sensitive_text(error_message, 500) if error_message else None,
             query_count, now, now),
        )


async def search_one_site(
    task_id: int, item: sqlite3.Row, site: dict[str, Any], clients: dict[str, Any], semaphore: asyncio.Semaphore,
    queries: list[tuple[str, str | None, str]],
) -> tuple[dict[str, Any], list[dict[str, Any]], str | None, int]:
    started = time.monotonic()
    async with semaphore:
        query_count = 0
        try:
            client = clients.get(str(site["adapter"]))
            if client is None:
                raise RuntimeError(f"不支持的站点适配器：{site['adapter']}")
            unique: dict[str, dict[str, Any]] = {}
            errors: list[Exception] = []
            site_queries = queries[:1] if str(site["adapter"]) == "rss" else queries
            for title, imdb_id, _label in site_queries:
                query_count += 1
                try:
                    rows = await client.search(site, title, imdb_id)
                except Exception as exc:
                    errors.append(exc)
                    continue
                for torrent in rows:
                    torrent = dict(torrent)
                    torrent["_site_priority"] = int(site.get("priority") or 100)
                    torrent["_site_id"] = int(site["id"])
                    key = str(torrent.get("enclosure") or resource_fingerprint(
                        str(first_value(torrent, ("title", "torrent_name", "name"), "")),
                        first_value(torrent, ("size", "size_bytes")),
                    ))
                    unique[key] = torrent
            torrents = list(unique.values())
            if not torrents and errors and len(errors) == query_count:
                raise errors[-1]
            duration_ms = max(0, int((time.monotonic() - started) * 1000))
            record_search_attempt(
                task_id, int(item["id"]), site, 1, "success", len(torrents), duration_ms,
                query_count=query_count,
            )
            return site, torrents, None, query_count
        except Exception as exc:
            error_code, reason = classify_search_error(exc)
            duration_ms = max(0, int((time.monotonic() - started) * 1000))
            record_search_attempt(
                task_id, int(item["id"]), site, 1, "failed", 0, duration_ms, error_code, safe_error(exc),
                query_count=max(1, query_count),
            )
            return site, [], reason, max(1, query_count)


async def run_search(task_id: int) -> None:
    emby, torznab, mteam, nexusphp, rss, config = EmbyClient(), TorznabClient(), MTeamClient(), NexusPHPClient(), RSSClient(), config_values()
    with connect() as conn:
        task = conn.execute("SELECT * FROM search_tasks WHERE id=?", (task_id,)).fetchone()
        if not task:
            return
        items = list(conn.execute(
            "SELECT * FROM playlist_items WHERE playlist_id=? AND rank_no BETWEEN ? AND ? ORDER BY rank_no",
            (task["playlist_id"], task["range_start"], task["range_end"]),
        ).fetchall())
        sites = rows_to_dicts(conn.execute("SELECT * FROM pt_sites WHERE enabled=1 AND search_enabled=1 ORDER BY priority,id").fetchall())
    try:
        selected_item_ids = {int(value) for value in json.loads(task["item_ids_json"] or "[]")}
        selected_site_ids = {int(value) for value in json.loads(task["site_ids_json"] or "[]")}
    except (TypeError, ValueError, json.JSONDecodeError):
        selected_item_ids, selected_site_ids = set(), set()
    if selected_item_ids:
        items = [item for item in items if int(item["id"]) in selected_item_ids]
    if selected_site_ids:
        sites = [site for site in sites if int(site["id"]) in selected_site_ids]
    update_task(task_id, status="running")
    task_log(task_id, "info", "task", f"开始搜索，共 {len(items)} 部影片、{len(sites)} 个搜索来源")
    matched = 0
    warnings: list[str] = []
    try:
        for completed, item in enumerate(items, start=1):
            label = f"#{item['rank_no']} {item['original_title']}"
            task_log(task_id, "info", "recognize", f"开始识别 {label}")
            try:
                tmdb_media = await recognize_movie(item["original_title"], item["year"], item["imdb_id"])
                if not tmdb_media:
                    raise RuntimeError("TMDB 未返回匹配结果")
                tmdb_id = int(tmdb_media["id"])
                media = {
                    "source": "themoviedb",
                    "tmdb_id": tmdb_id,
                    "imdb_id": tmdb_media.get("imdb_id") or item["imdb_id"],
                    "title": tmdb_media.get("title") or item["chinese_title"] or item["original_title"],
                    "original_title": tmdb_media.get("original_title") or item["original_title"],
                    "year": str(tmdb_media.get("release_date") or item["year"] or "")[:4] or None,
                    "release_date": tmdb_media.get("release_date"),
                    "type": "电影",
                    "poster_path": tmdb_media.get("poster_path"),
                }
                persist_tmdb_item(int(item["id"]), tmdb_media, item["imdb_id"])
                task_log(task_id, "info", "recognize", f"识别完成 {label} → TMDB {tmdb_id}")
                state, emby_item_id, image_tag = await library_details(
                    emby, str(media["title"]), int(media["year"]) if media.get("year") else item["year"],
                    tmdb_id, media.get("imdb_id"),
                )
                with connect() as conn:
                    conn.execute(
                        "UPDATE playlist_items SET library_state=?,library_checked_at=?,emby_item_id=?,emby_image_tag=? WHERE id=?",
                        (state, utc_now(), emby_item_id, image_tag, item["id"]),
                    )
                task_log(task_id, "info", "library", f"Emby 状态：{state}")
                if state == "in_library":
                    task_log(task_id, "info", "search", f"{label} 已有实体文件，跳过站点搜索")
                    update_task(task_id, completed=completed, matched=matched)
                    continue
                pairs: list[tuple[dict[str, Any] | None, dict[str, Any]]] = []
                clients = {"torznab": torznab, "mteam": mteam, "nexusphp": nexusphp, "rss": rss}
                site_semaphore = asyncio.Semaphore(4)
                queries = build_search_queries(item, media)
                task_log(task_id, "info", "search", "检索词：" + " → ".join(query[2] for query in queries))
                site_results = await asyncio.gather(*(
                    search_one_site(task_id, item, site, clients, site_semaphore, queries) for site in sites
                ))
                for site, torrents, reason, query_count in site_results:
                    if reason:
                        warnings.append(f"{label} · {site['name']}：{reason}")
                        task_log(task_id, "warning", "search", f"{site['name']} 搜索失败：{reason}")
                        continue
                    pairs.extend((media, torrent) for torrent in torrents)
                    task_log(task_id, "info", "search", f"{site['name']} 返回 {len(torrents)} 个资源（{query_count} 个检索词）")
                pairs.sort(key=lambda pair: analyze_candidate(str(first_value(pair[1], ("title", "torrent_name", "name"), "")), 0, config, pair[1])["ranking"])
                try:
                    policy = normalized_policy(json.loads(config.get("candidate_policy") or "{}"))
                except (TypeError, json.JSONDecodeError):
                    policy = normalized_policy({})
                limit = int(policy["candidate_limit"])
                eligible_keys: list[str] = []
                excluded_keys: list[str] = []
                selected_pairs: list[tuple[dict[str, Any] | None, dict[str, Any]]] = []
                selected_keys: set[tuple[str, str, str]] = set()
                for pair in pairs:
                    torrent = pair[1]
                    analysis = analyze_candidate(
                        str(first_value(torrent, ("title", "torrent_name", "name"), "")), 0, config, torrent,
                    )
                    identity_ok, identity_reason = candidate_identity(
                        item, media, str(first_value(torrent, ("title", "torrent_name", "name"), "")),
                    )
                    if not identity_ok:
                        analysis = dict(analysis)
                        analysis.update({
                            "eligible": False, "manual": True, "recommendation": "excluded",
                            "reason": identity_reason, "exclusion_reason": identity_reason,
                        })
                    key = resource_fingerprint(
                        str(first_value(torrent, ("title", "torrent_name", "name"), "")),
                        first_value(torrent, ("size", "size_bytes")),
                    )
                    bucket_name = "eligible" if analysis["eligible"] else "excluded"
                    bucket = eligible_keys if analysis["eligible"] else excluded_keys
                    if key not in bucket:
                        if len(bucket) >= limit:
                            continue
                        bucket.append(key)
                    selection_key = (bucket_name, key, str(first_value(torrent, ("site_name", "site"), torrent.get("_site_id") or "")))
                    if selection_key in selected_keys:
                        continue
                    selected_keys.add(selection_key)
                    selected_pairs.append(pair)
                for index, (source_media, torrent) in enumerate(selected_pairs):
                    candidate_id = uuid.uuid4().hex
                    title = str(first_value(torrent, ("title", "torrent_name", "name"), "未知资源"))
                    analyzed = analyze_candidate(title, index, config, torrent)
                    identity_ok, identity_reason = candidate_identity(item, source_media, title)
                    if not identity_ok:
                        analyzed = dict(analyzed)
                        analyzed.update({
                            "eligible": False, "manual": True, "recommendation": "excluded",
                            "reason": identity_reason, "exclusion_reason": identity_reason,
                        })
                    metadata = secret_free({
                        "description": torrent.get("description"), "labels": torrent.get("labels", []),
                        "volume_factor": torrent.get("volume_factor"), "publish_time": first_value(torrent, ("pubdate", "publish_time")),
                        "source": analyzed["source"], "profile_label": analyzed.get("profile_label"),
                    })
                    fingerprint = resource_fingerprint(title, first_value(torrent, ("size", "size_bytes")))
                    with connect() as conn:
                        conn.execute(
                            """INSERT INTO candidates(id,task_id,playlist_item_id,candidate_index,title,site_name,size,seeders,resolution,codec,group_name,group_tier,score,score_breakdown,ranking,recommendation,recommendation_reason,resource_key,library_state,is_manual_only,eligibility,exclusion_reason,profile_id,metadata_json,created_at)
                               VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                            (candidate_id, task_id, item["id"], index, title, first_value(torrent, ("site_name", "site")),
                             first_value(torrent, ("size", "size_bytes")), first_value(torrent, ("seeders", "seeder")), analyzed["resolution"],
                             analyzed["codec"], analyzed["group"], analyzed["tier"], analyzed["score"], json_value(analyzed["breakdown"]),
                             analyzed["ranking"], analyzed["recommendation"], analyzed["reason"], fingerprint, state,
                             int(analyzed["manual"]), "eligible" if analyzed["eligible"] else "excluded",
                             analyzed.get("exclusion_reason"), analyzed.get("profile_id"), json_value(metadata), utc_now()),
                        )
                    raw_candidates[candidate_id] = {"media": source_media, "torrent": torrent, "tmdb_id": tmdb_id}
                matched += len(eligible_keys)
                task_log(
                    task_id, "info", "candidate",
                    f"{label} 保留 {len(eligible_keys)} 个可下载候选，记录 {len(excluded_keys)} 个排除样本",
                )
            except Exception as exc:
                reason = safe_error(exc)
                warnings.append(f"{label}：{reason}")
                task_log(task_id, "error", "movie", f"{label} 处理失败：{reason}")
            update_task(task_id, completed=completed, matched=matched)
        status = "partial" if warnings else "completed"
        message = "；".join(warnings[:5])[:500] if warnings else None
        update_task(task_id, status=status, completed=len(items), matched=matched, error_message=message)
        task_log(task_id, "warning" if warnings else "info", "task", f"任务结束：{len(items)} 部已处理，{matched} 个候选，{len(warnings)} 个警告")
    except asyncio.CancelledError:
        update_task(task_id, status="cancelled")
        raise
    except Exception as exc:
        reason = safe_error(exc)
        update_task(task_id, status="failed", error_message=reason)
        task_log(task_id, "error", "task", f"任务异常停止：{reason}")
    finally:
        running_tasks.pop(task_id, None)


def update_recognition_task(task_id: int, **values: Any) -> None:
    if not set(values).issubset({"status", "completed", "matched", "error_message"}):
        raise ValueError("无效的识别任务字段")
    values["updated_at"] = utc_now()
    assignments = ", ".join(f"{key}=?" for key in values)
    with connect() as conn:
        # Dynamic column names are restricted by the allowlist above.
        conn.execute(f"UPDATE recognition_tasks SET {assignments} WHERE id=?", (*values.values(), task_id))  # nosec B608


async def run_recognition(task_id: int) -> None:
    with connect() as conn:
        task = conn.execute("SELECT * FROM recognition_tasks WHERE id=?", (task_id,)).fetchone()
        items = conn.execute(
            """SELECT * FROM playlist_items
               WHERE playlist_id=? AND (tmdb_id IS NULL OR tmdb_title IS NULL OR tmdb_original_title IS NULL)
               ORDER BY rank_no""", (task["playlist_id"],),
        ).fetchall()
    update_recognition_task(task_id, status="running")
    matched, errors = 0, []
    try:
        for completed, item in enumerate(items, start=1):
            try:
                media = await recognize_movie(item["original_title"], item["year"], item["imdb_id"])
                if media:
                    persist_tmdb_item(int(item["id"]), media, item["imdb_id"])
                    matched += 1
                else:
                    errors.append(f"#{item['rank_no']} 未识别")
            except Exception as exc:
                errors.append(f"#{item['rank_no']} {safe_error(exc)}")
            update_recognition_task(task_id, completed=completed, matched=matched)
        update_recognition_task(
            task_id, status="partial" if errors else "completed", completed=len(items), matched=matched,
            error_message="；".join(errors[:8])[:500] if errors else None,
        )
    except asyncio.CancelledError:
        update_recognition_task(task_id, status="cancelled")
        raise
    except Exception as exc:
        update_recognition_task(task_id, status="failed", error_message=safe_error(exc))
    finally:
        running_recognition_tasks.pop(task_id, None)


def update_automation_run(run_id: int, **values: Any) -> None:
    allowed = {"status", "stage", "total", "completed", "recognized", "searched", "recommended", "message"}
    if not set(values).issubset(allowed):
        raise ValueError("无效的自动化任务字段")
    values["updated_at"] = utc_now()
    assignments = ", ".join(f"{key}=?" for key in values)
    with connect() as conn:
        conn.execute(f"UPDATE automation_runs SET {assignments} WHERE id=?", (*values.values(), run_id))  # nosec B608


def add_notification(title: str, message: str, level: str = "info") -> None:
    with connect() as conn:
        conn.execute(
            "INSERT INTO notifications(level,title,message,created_at) VALUES(?,?,?,?)",
            (level, title[:120], sanitize_sensitive_text(message, 500), utc_now()),
        )


async def run_playlist_automation(run_id: int) -> None:
    try:
        with connect() as conn:
            run = conn.execute("SELECT * FROM automation_runs WHERE id=?", (run_id,)).fetchone()
            playlist = conn.execute("SELECT * FROM playlists WHERE id=?", (run["playlist_id"],)).fetchone() if run else None
        if not run or not playlist:
            return
        queue = await searchable_playlist_items(int(playlist["id"]), int(playlist["automation_batch_size"] or 50))
        items = queue["items"]
        update_automation_run(run_id, status="running", stage="recognition", total=len(items))
        emby = EmbyClient()
        recognized = searched = recommended = 0
        searchable_ids: list[int] = []
        for completed, item in enumerate(items, start=1):
            tmdb_id = item["tmdb_id"]
            if not tmdb_id or not item["tmdb_title"] or not item["tmdb_original_title"]:
                media = await recognize_movie(item["original_title"], item["year"], item["imdb_id"])
                if media:
                    tmdb_id = int(media["id"])
                    recognized += 1
                    persist_tmdb_item(int(item["id"]), media, item["imdb_id"])
                    item = dict(item)
                    item.update({
                        "tmdb_title": media.get("title"), "tmdb_original_title": media.get("original_title"),
                        "tmdb_year": str(media.get("release_date") or "")[:4] or item["year"],
                        "tmdb_imdb_id": media.get("imdb_id") or item["imdb_id"],
                    })
            update_automation_run(run_id, stage="library", completed=completed, recognized=recognized)
            state, emby_item_id, image_tag = await library_details(
                emby, canonical_item_title(item), canonical_item_year(item), tmdb_id,
                item["tmdb_imdb_id"] or item["imdb_id"],
            )
            with connect() as conn:
                conn.execute(
                    "UPDATE playlist_items SET library_state=?,library_checked_at=?,emby_item_id=?,emby_image_tag=? WHERE id=?",
                    (state, utc_now(), emby_item_id, image_tag, item["id"]),
                )
            if state != "in_library" and tmdb_id:
                searchable_ids.append(int(item["id"]))
        if searchable_ids:
            with connect() as conn:
                ranks = conn.execute(
                    f"SELECT MIN(rank_no),MAX(rank_no) FROM playlist_items WHERE id IN ({','.join('?' for _ in searchable_ids)})",  # nosec B608
                    searchable_ids,
                ).fetchone()
                now = utc_now()
                task_id = int(conn.execute(
                    """INSERT INTO search_tasks(playlist_id,range_start,range_end,status,total,trigger,item_ids_json,created_at,updated_at)
                       VALUES(?,?,?,?,?,?,?,?,?)""",
                    (playlist["id"], ranks[0], ranks[1], "queued", len(searchable_ids), "automation", json_value(searchable_ids), now, now),
                ).lastrowid)
            update_automation_run(run_id, stage="search")
            running_tasks[task_id] = asyncio.current_task()  # visible in health while the nested search runs
            await run_search(task_id)
            searched = len(searchable_ids)
            with connect() as conn:
                preferred_rows = conn.execute(
                    """SELECT c.id,c.playlist_item_id FROM candidates c
                       WHERE c.task_id=? AND c.recommendation='preferred' AND c.eligibility='eligible'
                       ORDER BY c.playlist_item_id,c.ranking""", (task_id,),
                ).fetchall()
                preferred = []
                seen_items: set[int] = set()
                for row in preferred_rows:
                    if int(row["playlist_item_id"]) not in seen_items:
                        preferred.append(row)
                        seen_items.add(int(row["playlist_item_id"]))
                recommended = len(preferred)
                if playlist["automation_auto_cart"]:
                    conn.executemany(
                        "INSERT OR IGNORE INTO cart_items(candidate_id,selected_at) VALUES(?,?)",
                        [(row["id"], utc_now()) for row in preferred],
                    )
        message = f"处理 {len(items)} 部，新增识别 {recognized}，搜索 {searched}，推荐 {recommended}；未自动下载"
        update_automation_run(
            run_id, status="completed", stage="completed", completed=len(items), recognized=recognized,
            searched=searched, recommended=recommended, message=message,
        )
        add_notification("新增影片处理完成", message, "success")
    except asyncio.CancelledError:
        update_automation_run(run_id, status="cancelled", message="新片处理任务已取消")
        raise
    except Exception as exc:
        reason = safe_error(exc)
        update_automation_run(run_id, status="failed", message=reason)
        add_notification("新增影片处理失败", reason, "error")
    finally:
        running_automation_tasks.pop(run_id, None)


async def sync_playlist_incremental(playlist_id: int, trigger: str = "manual") -> dict[str, Any]:
    with connect() as conn:
        playlist = conn.execute("SELECT * FROM playlists WHERE id=?", (playlist_id,)).fetchone()
    if not playlist:
        raise HTTPException(404, "片单不存在")
    if not playlist["source_url"]:
        raise HTTPException(422, "该片单没有可同步的网址来源")
    try:
        source = await PlaylistSourceFetcher().fetch(str(playlist["source_url"]), 10000)
        incoming = normalize_import_items(source.get("items", []))
        with connect() as conn:
            existing = conn.execute(
                "SELECT imdb_id,tmdb_id,original_title,year FROM playlist_items WHERE playlist_id=?", (playlist_id,),
            ).fetchall()
            keys = {
                ("imdb", str(row["imdb_id"])) if row["imdb_id"] else
                ("tmdb", str(row["tmdb_id"])) if row["tmdb_id"] else
                ("title", re.sub(r"\W+", "", str(row["original_title"]).lower()), str(row["year"] or ""))
                for row in existing
            }
            max_rank = int(conn.execute("SELECT COALESCE(MAX(rank_no),0) FROM playlist_items WHERE playlist_id=?", (playlist_id,)).fetchone()[0])
            additions = []
            for item in incoming:
                key = (("imdb", str(item["imdb_id"])) if item.get("imdb_id") else
                       ("tmdb", str(item["tmdb_id"])) if item.get("tmdb_id") else
                       ("title", re.sub(r"\W+", "", str(item["original_title"]).lower()), str(item.get("year") or "")))
                if key in keys:
                    continue
                keys.add(key)
                max_rank += 1
                additions.append({**item, "playlist_id": playlist_id, "rank_no": max_rank})
            if additions:
                conn.executemany(
                    """INSERT INTO playlist_items(playlist_id,rank_no,imdb_id,original_title,year,chinese_title,tmdb_id)
                       VALUES(:playlist_id,:rank_no,:imdb_id,:original_title,:year,:chinese_title,:tmdb_id)""", additions,
                )
            next_sync = (datetime.now(timezone.utc) + timedelta(hours=int(playlist["sync_interval_hours"] or 24))).isoformat()
            message = f"增量同步完成，新增 {len(additions)} 部，保留现有 {len(existing)} 部"
            conn.execute(
                """UPDATE playlists SET source_name=?,last_synced_at=?,next_sync_at=?,last_sync_status='completed',last_sync_message=?
                   WHERE id=?""", (source.get("source_name"), utc_now(), next_sync, message, playlist_id),
            )
        add_notification("片单来源已同步", f"{playlist['name']}：{message}", "success")
        if additions and playlist["automation_enabled"]:
            await start_playlist_automation(playlist_id, "sync")
        return {"id": playlist_id, "added": len(additions), "message": message, "trigger": trigger}
    except HTTPException:
        raise
    except Exception as exc:
        reason = safe_error(exc)
        with connect() as conn:
            conn.execute(
                "UPDATE playlists SET last_sync_status='failed',last_sync_message=?,next_sync_at=? WHERE id=?",
                (reason, (datetime.now(timezone.utc) + timedelta(hours=1)).isoformat(), playlist_id),
            )
        add_notification("片单同步失败", f"{playlist['name']}：{reason}", "error")
        raise


async def start_playlist_automation(playlist_id: int, trigger: str = "manual") -> dict[str, Any]:
    with connect() as conn:
        playlist = conn.execute("SELECT * FROM playlists WHERE id=?", (playlist_id,)).fetchone()
        if not playlist:
            raise HTTPException(404, "片单不存在")
        active = conn.execute(
            "SELECT id FROM automation_runs WHERE playlist_id=? AND status IN ('queued','running') ORDER BY id DESC LIMIT 1", (playlist_id,),
        ).fetchone()
        if active:
            return {"id": int(active["id"]), "status": "running", "message": "新增影片处理任务正在运行"}
        now = utc_now()
        run_id = int(conn.execute(
            "INSERT INTO automation_runs(playlist_id,trigger,status,stage,created_at,updated_at) VALUES(?,?,?,?,?,?)",
            (playlist_id, trigger, "queued", "queued", now, now),
        ).lastrowid)
    running_automation_tasks[run_id] = asyncio.create_task(run_playlist_automation(run_id))
    return {"id": run_id, "status": "queued", "message": "已开始识别并搜索未入库影片；不会自动下载"}


async def sync_scheduler() -> None:
    while True:
        await asyncio.sleep(60)
        cleanup_old_data()
        now = utc_now()
        with connect() as conn:
            due = [int(row["id"]) for row in conn.execute(
                """SELECT id FROM playlists WHERE sync_enabled=1 AND source_url IS NOT NULL AND source_url!=''
                   AND (next_sync_at IS NULL OR next_sync_at<=?)""", (now,),
            ).fetchall()]
        for playlist_id in due:
            try:
                await sync_playlist_incremental(playlist_id, "schedule")
            except Exception:
                continue


@app.on_event("startup")
async def startup() -> None:
    global scheduler_task
    initialize()
    load_runtime_settings()
    cleanup_old_data()
    with connect() as conn:
        conn.execute(
            "UPDATE search_tasks SET status='interrupted', updated_at=? WHERE status IN ('queued','running')",
            (utc_now(),),
        )
        conn.execute(
            "UPDATE recognition_tasks SET status='interrupted', updated_at=? WHERE status IN ('queued','running')",
            (utc_now(),),
        )
        conn.execute(
            "UPDATE library_scan_tasks SET status='interrupted', updated_at=? WHERE status IN ('queued','running')",
            (utc_now(),),
        )
        conn.execute(
            "UPDATE automation_runs SET status='interrupted', updated_at=? WHERE status IN ('queued','running')",
            (utc_now(),),
        )
    scheduler_task = asyncio.create_task(sync_scheduler())


@app.on_event("shutdown")
async def shutdown() -> None:
    if scheduler_task:
        scheduler_task.cancel()
        await asyncio.gather(scheduler_task, return_exceptions=True)


@app.get("/api/health")
async def health() -> dict[str, Any]:
    return {
        "ok": True,
        "version": app.version,
        "running_tasks": len(running_tasks) + len(running_automation_tasks),
        "access_token_required": access_token_required(),
    }


@app.get("/cookiecloud")
@app.get("/cookiecloud/")
async def cookiecloud_root() -> Response:
    return Response("AutoList CookieCloud API · /cookiecloud", media_type="text/plain")


@app.post("/cookiecloud/update")
async def cookiecloud_update(request: Request) -> dict[str, Any]:
    enforce_cookiecloud_rate_limit()
    try:
        content = decode_cookiecloud_body(await request.body(), request.headers.get("content-encoding", ""))
        payload = CookieCloudUploadPayload.model_validate_json(content)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, f"CookieCloud 上传数据无效：{safe_error(exc)}") from exc
    require_configured_cookiecloud_uuid(payload.uuid)
    path = cookiecloud_file(payload.uuid)
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(payload.model_dump(), ensure_ascii=False), encoding="utf-8")
    temporary.chmod(0o600)
    temporary.replace(path)
    path.chmod(0o600)
    forwarded = False
    if settings.cookiecloud_forward_moviepilot and settings.mp_base_url:
        try:
            forward_url = f"{settings.mp_base_url.rstrip('/')}/cookiecloud/update"
            async with httpx.AsyncClient(timeout=settings.mp_timeout_seconds) as client:
                response = await client.post(forward_url, json=payload.model_dump())
                forwarded = response.is_success
        except Exception:
            forwarded = False
    return {"action": "done", "forwarded_moviepilot": forwarded}


@app.get("/cookiecloud/get/{uuid_value}")
async def cookiecloud_get(uuid_value: str) -> dict[str, Any]:
    require_configured_cookiecloud_uuid(uuid_value)
    path = cookiecloud_file(uuid_value)
    if not path.exists():
        raise HTTPException(404, "CookieCloud 数据不存在")
    return json.loads(path.read_text(encoding="utf-8"))


@app.get("/api/cookiecloud/status")
async def cookiecloud_status() -> dict[str, Any]:
    configured = bool(settings.cookiecloud_key and settings.cookiecloud_password)
    path = cookiecloud_file(settings.cookiecloud_key) if settings.cookiecloud_key else None
    return {
        "configured": configured,
        "received": bool(path and path.exists()),
        "updated_at": datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc).isoformat() if path and path.exists() else None,
        "forward_moviepilot": settings.cookiecloud_forward_moviepilot,
        "endpoint": "/cookiecloud",
    }


@app.get("/api/connection")
async def connection() -> dict[str, Any]:
    async def check_one(name: str, client: Any) -> tuple[str, dict[str, Any]]:
        try:
            return name, await asyncio.wait_for(client.check(), timeout=6)
        except TimeoutError:
            return name, {"ok": False, "configured": True, "message": "连接检测超时"}
        except Exception as exc:
            return name, {"ok": False, "configured": True, "message": safe_error(exc)}
    checked = await asyncio.gather(*(
        check_one(name, client) for name, client in (
            ("tmdb", TMDBClient()), ("emby", EmbyClient()), ("transmission", TransmissionClient()), ("moviepilot", MoviePilotClient()),
        )
    ))
    results = dict(checked)
    # AutoList submits through MoviePilot; direct Transmission credentials are diagnostic only.
    required = (results["tmdb"], results["moviepilot"])
    return {"ok": all(item.get("ok") for item in required), "providers": results, "message": "核心服务正常" if all(item.get("ok") for item in required) else "核心服务需要配置"}


@app.get("/api/downloads")
async def downloads() -> list[dict[str, Any]]:
    try:
        return secret_free(await TransmissionClient().current_downloads())
    except Exception as exc:
        raise HTTPException(502, f"读取 Transmission 下载任务失败：{safe_error(exc)}") from exc


@app.get("/api/config")
async def get_config() -> dict[str, str]:
    return config_values()


def validated_base_url(value: str, label: str, required: bool) -> str:
    normalized = value.strip().rstrip("/")
    if not normalized:
        if required:
            raise HTTPException(422, f"{label}不能为空")
        return ""
    parsed = urlparse(normalized)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise HTTPException(422, f"{label}必须以 http:// 或 https:// 开头")
    if parsed.username or parsed.password:
        raise HTTPException(422, f"{label}不能包含用户名或密码")
    return normalized


@app.get("/api/settings")
async def get_runtime_settings() -> dict[str, Any]:
    return settings.public_values()


@app.put("/api/settings")
async def put_runtime_settings(payload: RuntimeSettingsPayload) -> dict[str, Any]:
    values = payload.model_dump()
    values["mp_base_url"] = validated_base_url(values["mp_base_url"], "MoviePilot 地址", False)
    values["emby_base_url"] = validated_base_url(values["emby_base_url"], "Emby 地址", False)
    values["ai_base_url"] = validated_base_url(values["ai_base_url"], "AI 地址", False)
    values["tr_base_url"] = validated_base_url(values["tr_base_url"], "Transmission 地址", False)
    if values.get("outbound_proxy_url"):
        values["outbound_proxy_url"] = validated_base_url(str(values["outbound_proxy_url"]), "代理地址", True)
    for key in ("mp_api_key", "emby_api_key", "tmdb_api_key", "mdblist_api_key", "cookiecloud_key", "cookiecloud_password", "ai_api_key", "tr_password"):
        if not values.get(key):
            values.pop(key, None)
    save_runtime_settings(values)
    return settings.public_values()


@app.post("/api/settings/test")
async def test_runtime_settings() -> dict[str, Any]:
    return (await connection())["providers"]


@app.get("/api/sites")
async def sites() -> list[dict[str, Any]]:
    with connect() as conn:
        rows = rows_to_dicts(conn.execute("SELECT * FROM pt_sites ORDER BY id").fetchall())
    try:
        mp_sites, statistics, users = await moviepilot_site_snapshot()
    except Exception:
        mp_sites, statistics, users = [], [], []
    for item in rows:
        item["api_key_configured"] = bool(item.get("api_key"))
        item["cookie_configured"] = bool(item.get("cookie"))
        item["api_key"] = ""
        item["cookie"] = ""
        item["rss_url_configured"] = bool(item.get("rss_url"))
        item["rss_url"] = ""
        mp_site = next((candidate for candidate in mp_sites if domain_match(str(candidate.get("domain") or candidate.get("url") or ""), item["base_url"])), None)
        stat = next((candidate for candidate in statistics if domain_match(str(candidate.get("domain") or ""), item["base_url"])), {})
        user = next((candidate for candidate in users if domain_match(str(candidate.get("domain") or ""), item["base_url"])), {})
        item["mp_site_id"] = mp_site.get("id") if mp_site else None
        if item["mp_site_id"]:
            moviepilot_site_ids[item["id"]] = int(item["mp_site_id"])
        item["icon_endpoint"] = f"/api/sites/{item['id']}/icon"
        item["mp_active"] = bool(mp_site.get("is_active")) if mp_site else bool(item["enabled"])
        item["mp_user"] = {key: user.get(key) for key in ("username", "user_level", "upload", "download", "ratio", "bonus", "seeding", "leeching", "updated_time", "err_msg")}
        if item["mp_user"].get("err_msg"):
            item["mp_user"]["err_msg"] = sanitize_sensitive_text(item["mp_user"]["err_msg"])
        item["mp_status"] = {key: stat.get(key) for key in ("seconds", "lst_state", "lst_mod_date", "success", "fail")}
    return rows


@app.get("/api/sites/{site_id}/icon")
async def site_icon(site_id: int) -> Response:
    """Serve MP's site icon locally so authenticated/private site favicons do not fail in the browser."""
    if site_id in site_icon_cache:
        content, media_type = site_icon_cache[site_id]
        return Response(content=content, media_type=media_type, headers={"Cache-Control": "public, max-age=86400"})
    with connect() as conn:
        row = conn.execute("SELECT id,base_url,icon_url FROM pt_sites WHERE id=?", (site_id,)).fetchone()
    if not row:
        raise HTTPException(404, "站点不存在")
    icon_value = str(row["icon_url"] or "")
    mp_id = moviepilot_site_ids.get(site_id)
    if not icon_value and not mp_id:
        try:
            mp_sites, _, _ = await moviepilot_site_snapshot()
            match = next((item for item in mp_sites if domain_match(str(item.get("domain") or item.get("url") or ""), row["base_url"])), None)
            mp_id = int(match["id"]) if match and match.get("id") else None
            if mp_id:
                moviepilot_site_ids[site_id] = mp_id
        except Exception:
            mp_id = None
    try:
        if mp_id:
            payload = await MoviePilotClient().site_icon(mp_id)
            data = payload.get("data", payload) if isinstance(payload, dict) else payload
            icon_value = str(data.get("icon") or data.get("url") or "") if isinstance(data, dict) else str(data or "")
        if icon_value.startswith("data:image/"):
            header, encoded = icon_value.split(",", 1)
            media_type = header.split(";", 1)[0].split(":", 1)[1]
            content = base64.b64decode(encoded)
        else:
            source = icon_value if icon_value.startswith(("http://", "https://")) else f"{str(row['base_url']).rstrip('/')}/favicon.ico"
            await validate_remote_icon_url(source, str(row["base_url"]))
            async with httpx.AsyncClient(timeout=12, follow_redirects=True) as client:
                upstream = await client.get(source)
                upstream.raise_for_status()
            await validate_remote_icon_url(str(upstream.url), str(row["base_url"]))
            if int(upstream.headers.get("content-length") or 0) > 2 * 1024 * 1024:
                raise RuntimeError("图标文件过大")
            content = upstream.content
            media_type = upstream.headers.get("content-type", "image/x-icon").split(";", 1)[0]
        detected_media_type = raster_image_media_type(content)
        if not content or len(content) > 2 * 1024 * 1024 or not detected_media_type:
            raise RuntimeError("图标内容无效")
        media_type = detected_media_type
        site_icon_cache[site_id] = (content, media_type)
        return Response(content=content, media_type=media_type, headers={"Cache-Control": "public, max-age=86400"})
    except Exception:
        # A deterministic SVG fallback still gives every site a consistent visual anchor.
        initial = (str(row["base_url"] or "?")[:1] or "?").upper()
        fallback = f'<svg xmlns="http://www.w3.org/2000/svg" width="64" height="64"><rect width="64" height="64" rx="14" fill="#eeeaff"/><text x="32" y="42" text-anchor="middle" font-family="Arial" font-size="28" font-weight="700" fill="#6657e8">{initial}</text></svg>'.encode()
        return Response(content=fallback, media_type="image/svg+xml", headers={"Cache-Control": "public, max-age=3600"})


@app.get("/api/sites/{site_id}/health-history")
async def site_health_history(site_id: int, limit: int = 50) -> dict[str, Any]:
    safe_limit = max(1, min(limit, 200))
    with connect() as conn:
        site = conn.execute("SELECT id,name FROM pt_sites WHERE id=?", (site_id,)).fetchone()
        if not site:
            raise HTTPException(404, "站点不存在")
        rows = conn.execute(
            """SELECT status,result_count,duration_ms,error_code,error_message,finished_at
               FROM search_attempts WHERE site_id=? ORDER BY id DESC LIMIT ?""", (site_id, safe_limit),
        ).fetchall()
        summary = conn.execute(
            """SELECT COUNT(*) AS total,SUM(CASE WHEN status='success' THEN 1 ELSE 0 END) AS succeeded,
                      CAST(AVG(duration_ms) AS INTEGER) AS average_ms,MAX(finished_at) AS last_attempt_at
               FROM search_attempts WHERE site_id=?""", (site_id,),
        ).fetchone()
    result = dict(summary)
    total = int(result["total"] or 0)
    result["success_rate"] = round(int(result["succeeded"] or 0) / total * 100, 1) if total else None
    return {"site": dict(site), "summary": result, "items": rows_to_dicts(rows)}


@app.post("/api/sites")
async def add_site(payload: SitePayload) -> dict[str, Any]:
    base_url = validated_base_url(payload.base_url, "站点地址", True)
    rss_url = validated_base_url(payload.rss_url, "RSS 地址", False)
    adapter = resolve_site_adapter(base_url, rss_url)
    try:
        with connect() as conn:
            site_id = conn.execute(
                """INSERT INTO pt_sites(name,adapter,base_url,api_key,cookie,user_agent,priority,timeout_seconds,rss_url,icon_url,proxy,render,limit_interval,limit_count,enabled,search_enabled,created_at)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (payload.name.strip(), adapter, base_url, payload.api_key or "", payload.cookie or "", payload.user_agent,
                 payload.priority, payload.timeout_seconds, rss_url, payload.icon_url, int(payload.proxy), int(payload.render),
                 payload.limit_interval, payload.limit_count, int(payload.enabled), int(payload.search_enabled), utc_now()),
            ).lastrowid
    except sqlite3.IntegrityError as exc:
        raise HTTPException(409, "站点名称已存在") from exc
    return {"id": site_id, "name": payload.name, "adapter": adapter, "enabled": payload.enabled, "search_enabled": payload.search_enabled}


@app.put("/api/sites/{site_id}")
async def update_site(site_id: int, payload: SitePayload) -> dict[str, Any]:
    base_url = validated_base_url(payload.base_url, "站点地址", True)
    with connect() as conn:
        current = conn.execute("SELECT * FROM pt_sites WHERE id=?", (site_id,)).fetchone()
        if not current:
            raise HTTPException(404, "站点不存在")
        rss_url = validated_base_url(payload.rss_url, "RSS 地址", False) if payload.rss_url.strip() else str(current["rss_url"] or "")
        adapter = resolve_site_adapter(base_url, rss_url)
        conn.execute(
            """UPDATE pt_sites SET name=?,adapter=?,base_url=?,api_key=?,cookie=?,user_agent=?,priority=?,timeout_seconds=?,rss_url=?,icon_url=?,proxy=?,render=?,limit_interval=?,limit_count=?,enabled=?,search_enabled=?,migration_note=NULL WHERE id=?""",
            (payload.name.strip(), adapter, base_url, payload.api_key or current["api_key"], payload.cookie or current["cookie"],
             payload.user_agent, payload.priority, payload.timeout_seconds, rss_url, payload.icon_url, int(payload.proxy),
             int(payload.render), payload.limit_interval, payload.limit_count, int(payload.enabled), int(payload.search_enabled), site_id),
        )
    site_icon_cache.pop(site_id, None)
    return {"id": site_id, "name": payload.name, "adapter": adapter, "enabled": payload.enabled, "search_enabled": payload.search_enabled}


@app.delete("/api/sites/{site_id}")
async def delete_site(site_id: int) -> dict[str, Any]:
    with connect() as conn:
        site = conn.execute("SELECT * FROM pt_sites WHERE id=?", (site_id,)).fetchone()
        if not site:
            raise HTTPException(404, "站点不存在")
        conn.execute("DELETE FROM pt_sites WHERE id=?", (site_id,))
    return {"deleted": site_id}


@app.post("/api/sites/{site_id}/test")
async def test_site(site_id: int) -> dict[str, Any]:
    with connect() as conn:
        row = conn.execute("SELECT * FROM pt_sites WHERE id=?", (site_id,)).fetchone()
    if not row:
        raise HTTPException(404, "站点不存在")
    return await test_site_config(dict(row))


@app.post("/api/sites/test")
async def test_all_sites() -> dict[str, Any]:
    with connect() as conn:
        rows = rows_to_dicts(conn.execute("SELECT * FROM pt_sites WHERE enabled=1 ORDER BY priority,id").fetchall())
    semaphore = asyncio.Semaphore(4)
    async def probe(site: dict[str, Any]) -> dict[str, Any]:
        async with semaphore:
            return await test_site_config(site)
    results = await asyncio.gather(*(probe(site) for site in rows))
    return {"total": len(results), "ok": sum(1 for result in results if result["ok"]), "results": results}


@app.post("/api/sites/sync-cookiecloud")
async def sync_sites_from_cookiecloud() -> dict[str, Any]:
    groups = cookie_groups(stored_cookiecloud_payload())
    updated: list[str] = []
    missing: list[str] = []
    with connect() as conn:
        rows = conn.execute("SELECT id,name,base_url FROM pt_sites ORDER BY id").fetchall()
        for row in rows:
            match = cookie_for_host(groups, urlparse(str(row["base_url"])).hostname or "")
            if not match:
                missing.append(str(row["name"]))
                continue
            conn.execute("UPDATE pt_sites SET cookie=?,migration_note=NULL WHERE id=?", (match[1], row["id"]))
            updated.append(str(row["name"]))
    return {
        "ok": bool(updated),
        "updated": len(updated),
        "sites": updated,
        "missing": missing,
        "message": f"已从 Chrome CookieCloud 更新 {len(updated)} 个站点 Cookie；UA 保留各站点现有配置",
    }


async def get_local_and_mp_site(site_id: int) -> tuple[dict[str, Any], dict[str, Any]]:
    with connect() as conn:
        row = conn.execute("SELECT * FROM pt_sites WHERE id=?", (site_id,)).fetchone()
    if not row:
        raise HTTPException(404, "站点不存在")
    local = dict(row)
    mp_sites, _, _ = await moviepilot_site_snapshot()
    mp_site = next((candidate for candidate in mp_sites if domain_match(str(candidate.get("domain") or candidate.get("url") or ""), local["base_url"])), None)
    if not mp_site:
        raise HTTPException(404, "MoviePilot 未找到对应站点")
    return local, mp_site


@app.post("/api/sites/{site_id}/sync-moviepilot")
async def sync_site_from_moviepilot(site_id: int) -> dict[str, Any]:
    local, mp_site = await get_local_and_mp_site(site_id)
    with connect() as conn:
        conn.execute("UPDATE pt_sites SET cookie=?,user_agent=?,proxy=?,render=?,migration_note=NULL WHERE id=?", (
            str(mp_site.get("cookie") or ""), str(mp_site.get("ua") or ""), int(bool(mp_site.get("proxy"))), int(bool(mp_site.get("render"))), local["id"],
        ))
    return {"ok": True, "name": local["name"], "message": "已从 MoviePilot 更新 Cookie 与 UA"}


@app.post("/api/sites/{site_id}/update-cookie-ua")
async def update_site_cookie_ua(site_id: int, payload: SiteCookiePayload) -> dict[str, Any]:
    local, mp_site = await get_local_and_mp_site(site_id)
    result = await MoviePilotClient().update_site_cookie(int(mp_site["id"]), payload.username, payload.password, payload.code)
    if not result.get("success", False):
        raise HTTPException(502, sanitize_sensitive_text(result.get("message") or "MoviePilot 更新 Cookie 失败"))
    await sync_site_from_moviepilot(site_id)
    return {"ok": True, "name": local["name"], "message": "MoviePilot 已更新并同步 Cookie 与 UA"}


@app.post("/api/sites/{site_id}/refresh-moviepilot-userdata")
async def refresh_site_moviepilot_userdata(site_id: int) -> dict[str, Any]:
    _, mp_site = await get_local_and_mp_site(site_id)
    result = await MoviePilotClient().refresh_site_user_data(int(mp_site["id"]))
    return {"ok": bool(result.get("success", True)), "message": sanitize_sensitive_text(result.get("message") or "已请求刷新用户数据")}


async def hydrate_recent_emby_posters(items: list[dict[str, Any]]) -> None:
    """Backfill legacy Emby references for the small home-page shelf without a full rescan."""
    if not items or not settings.emby_base_url or not settings.emby_api_key:
        return
    emby = EmbyClient()
    semaphore = asyncio.Semaphore(3)

    async def hydrate(item: dict[str, Any]) -> None:
        if item.get("emby_item_id") or item.get("library_state") not in {"in_library", "strm"}:
            return
        async with semaphore:
            _, emby_item_id, image_tag = await library_details(
                emby, item.get("tmdb_title") or item.get("chinese_title") or item["original_title"],
                item.get("tmdb_year") or item.get("year"), item.get("tmdb_id"),
                item.get("tmdb_imdb_id") or item.get("imdb_id"),
            )
        if not emby_item_id:
            return
        item["emby_item_id"] = emby_item_id
        item["emby_image_tag"] = image_tag
        with connect() as conn:
            conn.execute(
                "UPDATE playlist_items SET emby_item_id=?,emby_image_tag=? WHERE id=?",
                (emby_item_id, image_tag, item["id"]),
            )

    await asyncio.gather(*(hydrate(item) for item in items))


@app.get("/api/overview")
async def overview() -> dict[str, Any]:
    with connect() as conn:
        playlist = conn.execute(
            """SELECT p.id, p.name, COUNT(i.id) AS item_count FROM playlists p
               LEFT JOIN playlist_items i ON i.playlist_id=p.id GROUP BY p.id ORDER BY p.position,p.id LIMIT 1"""
        ).fetchone()
        playlist_stats = conn.execute(
            """SELECT
                 SUM(CASE WHEN tmdb_id IS NOT NULL THEN 1 ELSE 0 END) AS recognized_count,
                 SUM(CASE WHEN library_state='in_library' THEN 1 ELSE 0 END) AS in_library_count
               FROM playlist_items WHERE playlist_id=?""",
            (playlist["id"],),
        ).fetchone() if playlist else None
        if playlist and settings.dashboard_random_posters:
            daily_seed = int(datetime.now(timezone.utc).strftime("%Y%m%d"))
            recent_items = rows_to_dicts(conn.execute(
                """SELECT id,rank_no,imdb_id,tmdb_id,original_title,chinese_title,year,
                          tmdb_title,tmdb_original_title,tmdb_year,tmdb_imdb_id,
                          library_state,emby_item_id,emby_image_tag
                   FROM playlist_items WHERE playlist_id=? AND library_state='in_library'
                   ORDER BY ((id * (1103515245 + (? % 997))) & 2147483647),rank_no LIMIT 6""",
                (playlist["id"], daily_seed),
            ).fetchall())
        else:
            recent_items = rows_to_dicts(conn.execute(
                """SELECT id,rank_no,imdb_id,tmdb_id,original_title,chinese_title,year,
                          tmdb_title,tmdb_original_title,tmdb_year,tmdb_imdb_id,
                          library_state,emby_item_id,emby_image_tag
                   FROM playlist_items WHERE playlist_id=? ORDER BY rank_no LIMIT 6""",
                (playlist["id"],),
            ).fetchall()) if playlist else []
        latest_task = conn.execute("SELECT * FROM search_tasks ORDER BY id DESC LIMIT 1").fetchone()
        latest_candidates = conn.execute(
            "SELECT playlist_item_id,title,size,resource_key FROM candidates WHERE task_id=? AND eligibility='eligible'", (latest_task["id"],),
        ).fetchall() if latest_task else []
        latest_candidate_count = len({
            (row["playlist_item_id"], row["resource_key"] or resource_fingerprint(row["title"], row["size"]))
            for row in latest_candidates
        })
        cart_count = conn.execute("SELECT COUNT(*) FROM cart_items").fetchone()[0]
        history_count = conn.execute("SELECT COUNT(*) FROM download_history").fetchone()[0]
    await hydrate_recent_emby_posters(recent_items)
    for item in recent_items:
        item["poster_url"] = (
            f"/api/playlist-items/{item['id']}/poster?tag={item['emby_image_tag']}"
            if item.get("emby_item_id") and item.get("emby_image_tag") else None
        )
    return {
        "playlist_id": playlist["id"] if playlist else None,
        "playlist_name": playlist["name"] if playlist else None,
        "item_count": playlist["item_count"] if playlist else 0,
        "recognized_count": int(playlist_stats["recognized_count"] or 0) if playlist_stats else 0,
        "in_library_count": int(playlist_stats["in_library_count"] or 0) if playlist_stats else 0,
        "pending_count": max(0, int(playlist["item_count"] or 0) - int(playlist_stats["in_library_count"] or 0)) if playlist and playlist_stats else 0,
        "recent_items": recent_items,
        "cart_count": cart_count,
        "history_count": history_count,
        "latest_candidate_count": latest_candidate_count,
        "latest_task": dict(latest_task) if latest_task else None,
    }


@app.get("/api/playlist-items/{playlist_item_id}/poster")
async def playlist_item_poster(playlist_item_id: int, tag: str = "") -> Response:
    with connect() as conn:
        row = conn.execute(
            "SELECT emby_item_id,emby_image_tag FROM playlist_items WHERE id=?", (playlist_item_id,),
        ).fetchone()
    if not row or not row["emby_item_id"]:
        raise HTTPException(404, "影片没有可用的 Emby 海报")
    emby_item_id = str(row["emby_item_id"])
    if not re.fullmatch(r"[A-Za-z0-9-]{1,128}", emby_item_id):
        raise HTTPException(404, "Emby 影片标识无效")
    image_tag = str(row["emby_image_tag"] or tag or "")
    cache_key = f"{emby_item_id}:{image_tag}"
    if cache_key in poster_cache:
        content, media_type = poster_cache[cache_key]
    else:
        try:
            content, _ = await EmbyClient().poster(emby_item_id)
        except httpx.HTTPStatusError as exc:
            status = 404 if exc.response.status_code == 404 else 502
            raise HTTPException(status, "Emby 海报读取失败") from exc
        except Exception as exc:
            raise HTTPException(502, f"Emby 海报读取失败：{safe_error(exc)}") from exc
        media_type = raster_image_media_type(content) or ""
        if not media_type or len(content) > 8 * 1024 * 1024:
            raise HTTPException(422, "Emby 返回的海报格式无效")
        if len(poster_cache) >= 64:
            poster_cache.pop(next(iter(poster_cache)))
        poster_cache[cache_key] = (content, media_type)
    return Response(
        content=content, media_type=media_type,
        headers={"Cache-Control": "private, max-age=86400", "X-Content-Type-Options": "nosniff"},
    )


@app.put("/api/config")
async def put_config(payload: ConfigPayload) -> dict[str, str]:
    raw = payload.model_dump(exclude_none=True)
    if "candidate_policy" in raw:
        try:
            raw["candidate_policy"] = normalized_policy(raw["candidate_policy"])
        except (TypeError, ValueError) as exc:
            raise HTTPException(422, safe_error(exc)) from exc
        raw["candidate_limit"] = raw["candidate_policy"]["candidate_limit"]
    values = {
        key: json_value(value) if key in {"scoring_policy", "candidate_policy"} else str(value)
        for key, value in raw.items()
    }
    save_config(values)
    return config_values()


@app.post("/api/config/score-preview")
async def score_preview(payload: ScorePreviewPayload) -> dict[str, Any]:
    config = config_values()
    if payload.candidate_policy is not None:
        try:
            config["candidate_policy"] = json_value(normalized_policy(payload.candidate_policy))
        except (TypeError, ValueError) as exc:
            raise HTTPException(422, safe_error(exc)) from exc
    return analyze_candidate(payload.title, 0, config, {"seeders": payload.seeders, "volume_factor": payload.volume_factor})


@app.get("/api/config/release-groups")
async def release_groups() -> dict[str, Any]:
    config = config_values()
    try:
        policy = normalized_policy(json.loads(config.get("candidate_policy") or "{}"))
    except (TypeError, json.JSONDecodeError):
        policy = normalized_policy({})
    return release_group_catalog(policy)


@app.post("/api/config/release-groups/import-moviepilot")
async def import_moviepilot_release_groups() -> dict[str, Any]:
    """Import MP custom groups once, then keep the merged vocabulary inside AutoList."""
    try:
        imported = await MoviePilotClient().custom_release_groups()
    except Exception as exc:
        raise HTTPException(502, f"MoviePilot 自定义制作组读取失败：{safe_error(exc)}") from exc
    config = config_values()
    try:
        policy = normalized_policy(json.loads(config.get("candidate_policy") or "{}"))
    except (TypeError, json.JSONDecodeError):
        policy = normalized_policy({})
    before = len(policy["custom_release_groups"])
    try:
        policy["custom_release_groups"] = merge_custom_rules([*policy["custom_release_groups"], *imported])
    except ValueError as exc:
        raise HTTPException(422, safe_error(exc)) from exc
    save_config({"candidate_policy": json_value(policy), "candidate_limit": str(policy["candidate_limit"])})
    catalog = release_group_catalog(policy)
    return {**catalog, "imported": len(policy["custom_release_groups"]) - before,
            "message": f"已合并 {len(policy['custom_release_groups']) - before} 条 MoviePilot 自定义制作组规则"}


@app.post("/api/playlists/import/preview")
async def preview_playlist_import(payload: ImportPayload) -> dict[str, Any]:
    try:
        source_name, items, source = await resolve_import(payload)
    except (ValueError, httpx.HTTPError) as exc:
        raise HTTPException(422, f"片单获取失败：{safe_error(exc)}") from exc
    if not items:
        raise HTTPException(422, "没有获取到可导入的电影")
    return {"name": payload.name or source_name, "count": len(items), "source_type": source.get("source_type"), "sample": items[:8]}


@app.post("/api/playlists/import")
async def import_playlist(payload: ImportPayload) -> dict[str, Any]:
    try:
        source_name, items, source = await resolve_import(payload)
    except (ValueError, httpx.HTTPError) as exc:
        raise HTTPException(422, f"片单获取失败：{safe_error(exc)}") from exc
    if not items:
        raise HTTPException(422, "没有可导入的影片")
    name = payload.name or source_name or f"片单 {datetime.now().strftime('%Y-%m-%d')}"
    with connect() as conn:
        position = conn.execute("SELECT COALESCE(MAX(position),0)+1 FROM playlists").fetchone()[0]
        cursor = conn.execute(
            "INSERT INTO playlists(name,position,source_type,source_url,source_name,last_synced_at,created_at) VALUES(?,?,?,?,?,?,?)",
            (name, position, source.get("source_type"), source.get("source_url"), source.get("source_name"), utc_now(), utc_now()),
        )
        playlist_id = cursor.lastrowid
        conn.executemany(
            """INSERT INTO playlist_items(playlist_id,rank_no,imdb_id,original_title,year,chinese_title,tmdb_id)
               VALUES(:playlist_id,:rank_no,:imdb_id,:original_title,:year,:chinese_title,:tmdb_id)""",
            [{"playlist_id": playlist_id, "tmdb_id": None, **item} for item in items],
        )
    return {"id": playlist_id, "name": name, "count": len(items)}


@app.post("/api/playlists/{playlist_id}/refresh-source")
async def refresh_playlist_source(playlist_id: int) -> dict[str, Any]:
    with connect() as conn:
        playlist = conn.execute("SELECT * FROM playlists WHERE id=?", (playlist_id,)).fetchone()
        if not playlist:
            raise HTTPException(404, "片单不存在")
        if not playlist["source_url"]:
            raise HTTPException(422, "该片单不是通过网址导入的")
        active_tables = (
            ("search_tasks", "搜索"),
            ("recognition_tasks", "识别"),
            ("library_scan_tasks", "入库检查"),
        )
        for table, label in active_tables:
            if conn.execute(f"SELECT 1 FROM {table} WHERE playlist_id=? AND status IN ('queued','running')", (playlist_id,)).fetchone():  # nosec B608
                raise HTTPException(409, f"片单仍有{label}任务运行，请完成后再刷新")
    try:
        source = await PlaylistSourceFetcher().fetch(str(playlist["source_url"]), 10000)
    except (ValueError, httpx.HTTPError) as exc:
        raise HTTPException(422, f"片单刷新失败：{safe_error(exc)}") from exc
    items = normalize_import_items(source.get("items", []))
    if not items:
        raise HTTPException(422, "来源没有返回可用电影，已保留现有片单")
    with connect() as conn:
        conn.execute("DELETE FROM playlist_items WHERE playlist_id=?", (playlist_id,))
        conn.executemany(
            """INSERT INTO playlist_items(playlist_id,rank_no,imdb_id,original_title,year,chinese_title,tmdb_id)
               VALUES(:playlist_id,:rank_no,:imdb_id,:original_title,:year,:chinese_title,:tmdb_id)""",
            [{"playlist_id": playlist_id, **item} for item in items],
        )
        conn.execute("UPDATE playlists SET source_name=?,last_synced_at=? WHERE id=?", (source.get("source_name"), utc_now(), playlist_id))
    return {"id": playlist_id, "count": len(items), "message": f"已从来源刷新 {len(items)} 部电影"}


@app.get("/api/playlists")
async def playlists() -> list[dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute(
            """SELECT p.*, COUNT(i.id) AS item_count, SUM(CASE WHEN i.tmdb_id IS NOT NULL THEN 1 ELSE 0 END) AS recognized_count FROM playlists p
               LEFT JOIN playlist_items i ON i.playlist_id=p.id GROUP BY p.id ORDER BY p.position,p.id"""
        ).fetchall()
    return rows_to_dicts(rows)


@app.put("/api/playlists/{playlist_id}/automation")
async def configure_playlist_automation(playlist_id: int, payload: PlaylistAutomationPayload) -> dict[str, Any]:
    with connect() as conn:
        if not conn.execute("SELECT 1 FROM playlists WHERE id=?", (playlist_id,)).fetchone():
            raise HTTPException(404, "片单不存在")
        conn.execute(
            "UPDATE playlists SET automation_enabled=?,automation_auto_cart=?,automation_batch_size=? WHERE id=?",
            (int(payload.enabled), int(payload.auto_cart), payload.batch_size, playlist_id),
        )
    return {"id": playlist_id, **payload.model_dump(), "auto_download": False}


@app.post("/api/playlists/{playlist_id}/automation/run")
async def run_playlist_automation_now(playlist_id: int) -> dict[str, Any]:
    return await start_playlist_automation(playlist_id)


@app.get("/api/automation-runs")
async def automation_runs(playlist_id: int | None = None, limit: int = 30) -> list[dict[str, Any]]:
    safe_limit = max(1, min(limit, 100))
    with connect() as conn:
        if playlist_id is None:
            rows = conn.execute("SELECT * FROM automation_runs ORDER BY id DESC LIMIT ?", (safe_limit,)).fetchall()
        else:
            rows = conn.execute(
                "SELECT * FROM automation_runs WHERE playlist_id=? ORDER BY id DESC LIMIT ?", (playlist_id, safe_limit),
            ).fetchall()
    return rows_to_dicts(rows)


@app.put("/api/playlists/{playlist_id}/sync-settings")
async def configure_playlist_sync(playlist_id: int, payload: PlaylistSyncPayload) -> dict[str, Any]:
    with connect() as conn:
        playlist = conn.execute("SELECT source_url FROM playlists WHERE id=?", (playlist_id,)).fetchone()
        if not playlist:
            raise HTTPException(404, "片单不存在")
        if payload.enabled and not playlist["source_url"]:
            raise HTTPException(422, "只有网址导入的片单可以启用定时同步")
        next_sync = (datetime.now(timezone.utc) + timedelta(hours=payload.interval_hours)).isoformat() if payload.enabled else None
        conn.execute(
            "UPDATE playlists SET sync_enabled=?,sync_interval_hours=?,next_sync_at=? WHERE id=?",
            (int(payload.enabled), payload.interval_hours, next_sync, playlist_id),
        )
    return {"id": playlist_id, **payload.model_dump(), "next_sync_at": next_sync}


@app.post("/api/playlists/{playlist_id}/sync-now")
async def sync_playlist_now(playlist_id: int) -> dict[str, Any]:
    return await sync_playlist_incremental(playlist_id)


@app.get("/api/notifications")
async def notifications(limit: int = 30, unread_only: bool = False) -> list[dict[str, Any]]:
    safe_limit = max(1, min(limit, 100))
    with connect() as conn:
        rows = conn.execute(
            f"SELECT * FROM notifications {'WHERE read=0' if unread_only else ''} ORDER BY id DESC LIMIT ?",  # nosec B608
            (safe_limit,),
        ).fetchall()
    return rows_to_dicts(rows)


@app.post("/api/notifications/read-all")
async def read_all_notifications() -> dict[str, Any]:
    with connect() as conn:
        count = conn.execute("UPDATE notifications SET read=1 WHERE read=0").rowcount
    return {"updated": count}


@app.get("/api/playlists/{playlist_id}/items")
async def playlist_items(
    playlist_id: int,
    page: int | None = None,
    page_size: int = 50,
    query: str = "",
    library_state: str = "all",
) -> Any:
    """Return a backwards-compatible full list or a bounded, filtered page."""
    with connect() as conn:
        if not conn.execute("SELECT 1 FROM playlists WHERE id=?", (playlist_id,)).fetchone():
            raise HTTPException(404, "片单不存在")
        if page is None:
            rows = conn.execute("SELECT * FROM playlist_items WHERE playlist_id=? ORDER BY rank_no", (playlist_id,)).fetchall()
            return rows_to_dicts(rows)
        safe_page = max(1, page)
        safe_page_size = max(1, min(page_size, 200))
        conditions = ["playlist_id=?"]
        params: list[Any] = [playlist_id]
        normalized_query = query.strip()
        if normalized_query:
            conditions.append("(original_title LIKE ? OR chinese_title LIKE ? OR tmdb_title LIKE ? OR tmdb_original_title LIKE ? OR imdb_id LIKE ? OR tmdb_imdb_id LIKE ?)")
            pattern = f"%{normalized_query}%"
            params.extend((pattern, pattern, pattern, pattern, pattern, pattern))
        if library_state == "not_downloaded":
            conditions.append("library_state!='in_library'")
        elif library_state in {"in_library", "unknown"}:
            conditions.append("library_state=?")
            params.append(library_state)
        elif library_state != "all":
            raise HTTPException(422, "无效的 Emby 状态筛选")
        where = " AND ".join(conditions)
        # SQL fragments in `conditions` are fixed literals; user values remain bound parameters.
        total = int(conn.execute(f"SELECT COUNT(*) FROM playlist_items WHERE {where}", params).fetchone()[0])  # nosec B608
        pages = max(1, math.ceil(total / safe_page_size))
        safe_page = min(safe_page, pages)
        rows = conn.execute(
            f"SELECT * FROM playlist_items WHERE {where} ORDER BY rank_no LIMIT ? OFFSET ?",  # nosec B608
            (*params, safe_page_size, (safe_page - 1) * safe_page_size),
        ).fetchall()
    return {
        "items": rows_to_dicts(rows), "total": total, "page": safe_page,
        "page_size": safe_page_size, "pages": pages,
    }


@app.post("/api/playlists/{playlist_id}/library-scan")
async def scan_playlist_library(playlist_id: int) -> dict[str, Any]:
    with connect() as conn:
        if not conn.execute("SELECT 1 FROM playlists WHERE id=?", (playlist_id,)).fetchone():
            raise HTTPException(404, "片单不存在")
        active = conn.execute(
            "SELECT id FROM library_scan_tasks WHERE playlist_id=? AND status IN ('queued','running') ORDER BY id DESC LIMIT 1", (playlist_id,),
        ).fetchone()
        if active:
            return {"id": active["id"], "message": "Emby 状态刷新正在进行"}
        total = conn.execute("SELECT COUNT(*) FROM playlist_items WHERE playlist_id=?", (playlist_id,)).fetchone()[0]
        task_id = conn.execute(
            """INSERT INTO library_scan_tasks(playlist_id,status,total,created_at,updated_at)
               VALUES(?,?,?,?,?)""", (playlist_id, "queued", total, utc_now(), utc_now()),
        ).lastrowid
    running_library_tasks[task_id] = asyncio.create_task(run_library_scan(task_id))
    return {"id": task_id, "message": f"开始刷新 {total} 部影片的 Emby 状态"}


@app.get("/api/library-scan-tasks/{task_id}")
async def get_library_scan_task(task_id: int) -> dict[str, Any]:
    with connect() as conn:
        task = conn.execute("SELECT * FROM library_scan_tasks WHERE id=?", (task_id,)).fetchone()
    if not task:
        raise HTTPException(404, "状态刷新任务不存在")
    return dict(task)


@app.put("/api/playlists/{playlist_id}")
async def update_playlist(playlist_id: int, payload: PlaylistUpdatePayload) -> dict[str, Any]:
    with connect() as conn:
        if not conn.execute("SELECT 1 FROM playlists WHERE id=?", (playlist_id,)).fetchone():
            raise HTTPException(404, "片单不存在")
        conn.execute("UPDATE playlists SET name=? WHERE id=?", (payload.name.strip(), playlist_id))
    return {"id": playlist_id, "name": payload.name.strip()}


@app.post("/api/playlists/reorder")
async def reorder_playlists(payload: PlaylistOrderPayload) -> dict[str, Any]:
    with connect() as conn:
        existing = {row[0] for row in conn.execute("SELECT id FROM playlists")}
        if len(payload.ids) != len(existing) or len(set(payload.ids)) != len(payload.ids) or set(payload.ids) != existing:
            raise HTTPException(422, "排序列表必须包含全部片单")
        conn.executemany("UPDATE playlists SET position=? WHERE id=?", [(index, site_id) for index, site_id in enumerate(payload.ids, start=1)])
    return {"ids": payload.ids}


@app.post("/api/playlists/{playlist_id}/recognize")
async def recognize_playlist(playlist_id: int) -> dict[str, Any]:
    with connect() as conn:
        if not conn.execute("SELECT 1 FROM playlists WHERE id=?", (playlist_id,)).fetchone():
            raise HTTPException(404, "片单不存在")
        total = conn.execute(
            """SELECT COUNT(*) FROM playlist_items
               WHERE playlist_id=? AND (tmdb_id IS NULL OR tmdb_title IS NULL OR tmdb_original_title IS NULL)""",
            (playlist_id,),
        ).fetchone()[0]
        if not total:
            return {"id": None, "status": "completed", "total": 0, "message": "片单已全部识别"}
        active = conn.execute(
            "SELECT id FROM recognition_tasks WHERE playlist_id=? AND status IN ('queued','running') ORDER BY id DESC LIMIT 1", (playlist_id,),
        ).fetchone()
        if active:
            return {"id": active["id"], "status": "running", "total": total}
        task_id = conn.execute(
            "INSERT INTO recognition_tasks(playlist_id,status,total,created_at,updated_at) VALUES(?,?,?,?,?)",
            (playlist_id, "queued", total, utc_now(), utc_now()),
        ).lastrowid
    running_recognition_tasks[task_id] = asyncio.create_task(run_recognition(task_id))
    return {"id": task_id, "status": "queued", "total": total}


@app.get("/api/recognition-tasks/{task_id}")
async def recognition_task_status(task_id: int) -> dict[str, Any]:
    with connect() as conn:
        task = conn.execute("SELECT * FROM recognition_tasks WHERE id=?", (task_id,)).fetchone()
    if not task:
        raise HTTPException(404, "识别任务不存在")
    return dict(task)


@app.delete("/api/playlists/{playlist_id}")
async def delete_playlist(playlist_id: int) -> dict[str, Any]:
    tasks_to_cancel: list[asyncio.Task[None]] = []
    with connect() as conn:
        exists = conn.execute("SELECT 1 FROM playlists WHERE id=?", (playlist_id,)).fetchone()
        if not exists:
            raise HTTPException(404, "片单不存在")
        for table, registry in (
            ("search_tasks", running_tasks),
            ("recognition_tasks", running_recognition_tasks),
            ("library_scan_tasks", running_library_tasks),
        ):
            ids = conn.execute(
                f"SELECT id FROM {table} WHERE playlist_id=? AND status IN ('queued','running')", (playlist_id,),  # nosec B608
            ).fetchall()
            for row in ids:
                task = registry.get(int(row["id"]))
                if task and not task.done():
                    task.cancel()
                    tasks_to_cancel.append(task)
    if tasks_to_cancel:
        await asyncio.gather(*tasks_to_cancel, return_exceptions=True)
    with connect() as conn:
        conn.execute("DELETE FROM playlists WHERE id=?", (playlist_id,))
    return {"deleted": playlist_id}


@app.post("/api/search-tasks")
async def create_task(payload: TaskPayload) -> dict[str, Any]:
    enforce_search_task_capacity()
    if payload.scope == "range" and payload.range_end < payload.range_start:
        raise HTTPException(422, "结束序号不能小于起始序号")
    item_ids: list[int] = []
    if payload.scope == "pending":
        queue = await searchable_playlist_items(payload.playlist_id, payload.count)
        item_ids = [int(item["id"]) for item in queue["items"]]
        if not item_ids:
            raise HTTPException(422, "当前片单没有可搜索的未入库影片")
        range_start = min(int(item["rank_no"]) for item in queue["items"])
        range_end = max(int(item["rank_no"]) for item in queue["items"])
    else:
        range_start, range_end = payload.range_start, payload.range_end
    with connect() as conn:
        total = conn.execute(
            "SELECT COUNT(*) FROM playlist_items WHERE playlist_id=? AND rank_no BETWEEN ? AND ?",
            (payload.playlist_id, range_start, range_end),
        ).fetchone()[0]
        if not total:
            raise HTTPException(422, "所选范围没有影片")
        if not conn.execute("SELECT 1 FROM pt_sites WHERE enabled=1 AND search_enabled=1 LIMIT 1").fetchone():
            raise HTTPException(422, "请先在站点配置中选择至少一个参与搜索的站点")
        task_id = conn.execute(
            """INSERT INTO search_tasks(playlist_id,range_start,range_end,status,total,trigger,item_ids_json,created_at,updated_at)
               VALUES(?,?,?,?,?,?,?,?,?)""",
            (payload.playlist_id, range_start, range_end, "queued", len(item_ids) or total,
             "pending" if payload.scope == "pending" else "manual", json_value(item_ids) if item_ids else None,
             utc_now(), utc_now()),
        ).lastrowid
    running_tasks[task_id] = asyncio.create_task(run_search(task_id))
    return {"id": task_id, "status": "queued", "total": len(item_ids) or total, "scope": payload.scope}


@app.get("/api/playlists/{playlist_id}/searchable-items")
async def searchable_items(playlist_id: int, limit: int = 2000) -> dict[str, Any]:
    safe_limit = max(1, min(limit, 10000))
    return await searchable_playlist_items(playlist_id, safe_limit)


@app.get("/api/search-tasks")
async def search_tasks(playlist_id: int | None = None, limit: int = 20) -> list[dict[str, Any]]:
    safe_limit = max(1, min(limit, 100))
    with connect() as conn:
        if playlist_id is None:
            rows = conn.execute("SELECT * FROM search_tasks ORDER BY id DESC LIMIT ?", (safe_limit,)).fetchall()
        else:
            rows = conn.execute(
                "SELECT * FROM search_tasks WHERE playlist_id=? ORDER BY id DESC LIMIT ?",
                (playlist_id, safe_limit),
            ).fetchall()
        result = rows_to_dicts(rows)
        for task in result:
            summary = conn.execute(
                """SELECT COUNT(*) AS total,
                          SUM(CASE WHEN status='success' THEN 1 ELSE 0 END) AS succeeded,
                          SUM(CASE WHEN status='failed' THEN 1 ELSE 0 END) AS failed
                   FROM search_attempts WHERE task_id=?""", (task["id"],),
            ).fetchone()
            task["attempt_summary"] = dict(summary)
    return result


@app.post("/api/search-tasks/{task_id}/cancel")
async def cancel_task(task_id: int) -> dict[str, Any]:
    task = running_tasks.get(task_id)
    if task:
        task.cancel()
    with connect() as conn:
        exists = conn.execute("SELECT 1 FROM search_tasks WHERE id=?", (task_id,)).fetchone()
    if not exists:
        raise HTTPException(404, "搜索任务不存在")
    update_task(task_id, status="cancelled")
    return {"id": task_id, "status": "cancelled"}


@app.get("/api/search-tasks/{task_id}")
async def task_status(task_id: int) -> dict[str, Any]:
    with connect() as conn:
        task = conn.execute("SELECT * FROM search_tasks WHERE id=?", (task_id,)).fetchone()
        summary = conn.execute(
            """SELECT COUNT(*) AS total,
                      SUM(CASE WHEN status='success' THEN 1 ELSE 0 END) AS succeeded,
                      SUM(CASE WHEN status='failed' THEN 1 ELSE 0 END) AS failed
               FROM search_attempts WHERE task_id=?""", (task_id,),
        ).fetchone()
    if not task:
        raise HTTPException(404, "搜索任务不存在")
    result = dict(task)
    result["attempt_summary"] = dict(summary)
    return result


@app.get("/api/search-tasks/{task_id}/attempts")
async def task_attempts(task_id: int, limit: int = 500) -> dict[str, Any]:
    safe_limit = max(1, min(limit, 2000))
    with connect() as conn:
        if not conn.execute("SELECT 1 FROM search_tasks WHERE id=?", (task_id,)).fetchone():
            raise HTTPException(404, "搜索任务不存在")
        rows = conn.execute(
            """SELECT a.id,a.playlist_item_id,p.rank_no,p.original_title,a.site_id,a.site_name,
                      a.attempt_no,a.status,a.result_count,a.duration_ms,a.error_code,a.error_message,a.finished_at
               FROM search_attempts a JOIN playlist_items p ON p.id=a.playlist_item_id
               WHERE a.task_id=? ORDER BY a.id DESC LIMIT ?""", (task_id, safe_limit),
        ).fetchall()
        summary_rows = conn.execute(
            """SELECT site_id,site_name,COUNT(*) AS total,
                      SUM(CASE WHEN status='success' THEN 1 ELSE 0 END) AS succeeded,
                      SUM(CASE WHEN status='failed' THEN 1 ELSE 0 END) AS failed,
                      CAST(AVG(duration_ms) AS INTEGER) AS average_ms
               FROM search_attempts WHERE task_id=? GROUP BY site_id,site_name ORDER BY site_name""", (task_id,),
        ).fetchall()
    return {"items": list(reversed(rows_to_dicts(rows))), "sites": rows_to_dicts(summary_rows)}


def create_followup_search_task(task_id: int, failed_only: bool) -> tuple[int, int]:
    with connect() as conn:
        source = conn.execute("SELECT * FROM search_tasks WHERE id=?", (task_id,)).fetchone()
        if not source:
            raise HTTPException(404, "搜索任务不存在")
        site_ids: list[int] = []
        item_ids: list[int] = []
        if failed_only:
            failed = conn.execute(
                """SELECT DISTINCT site_id,playlist_item_id FROM search_attempts
                   WHERE task_id=? AND status='failed' AND site_id IS NOT NULL""", (task_id,),
            ).fetchall()
            site_ids = sorted({int(row["site_id"]) for row in failed})
            item_ids = sorted({int(row["playlist_item_id"]) for row in failed})
            if not failed:
                raise HTTPException(422, "该任务没有可重试的站点失败记录")
        total = len(item_ids) if failed_only else int(source["total"])
        now = utc_now()
        new_id = conn.execute(
            """INSERT INTO search_tasks(
                 playlist_id,range_start,range_end,status,total,parent_task_id,trigger,site_ids_json,item_ids_json,created_at,updated_at
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
            (source["playlist_id"], source["range_start"], source["range_end"], "queued", total, task_id,
             "retry" if failed_only else "restart", json_value(site_ids) if site_ids else None,
             json_value(item_ids) if item_ids else None, now, now),
        ).lastrowid
    return int(new_id), total


@app.post("/api/search-tasks/{task_id}/retry")
async def retry_task(task_id: int) -> dict[str, Any]:
    enforce_search_task_capacity()
    new_id, total = create_followup_search_task(task_id, True)
    running_tasks[new_id] = asyncio.create_task(run_search(new_id))
    return {"id": new_id, "status": "queued", "total": total, "parent_task_id": task_id}


@app.post("/api/search-tasks/{task_id}/restart")
async def restart_task(task_id: int) -> dict[str, Any]:
    enforce_search_task_capacity()
    with connect() as conn:
        source = conn.execute("SELECT status FROM search_tasks WHERE id=?", (task_id,)).fetchone()
    if not source:
        raise HTTPException(404, "搜索任务不存在")
    if source["status"] in {"queued", "running"}:
        raise HTTPException(409, "任务仍在执行，无需重新启动")
    new_id, total = create_followup_search_task(task_id, False)
    running_tasks[new_id] = asyncio.create_task(run_search(new_id))
    return {"id": new_id, "status": "queued", "total": total, "parent_task_id": task_id}


@app.get("/api/search-tasks/{task_id}/logs")
async def task_logs(task_id: int, limit: int = 200) -> list[dict[str, Any]]:
    safe_limit = max(1, min(limit, 500))
    with connect() as conn:
        if not conn.execute("SELECT 1 FROM search_tasks WHERE id=?", (task_id,)).fetchone():
            raise HTTPException(404, "搜索任务不存在")
        rows = conn.execute(
            "SELECT id,level,stage,message,created_at FROM search_task_logs WHERE task_id=? ORDER BY id DESC LIMIT ?",
            (task_id, safe_limit),
        ).fetchall()
    return list(reversed(rows_to_dicts(rows)))


@app.get("/api/candidates")
async def candidates(task_id: int) -> list[dict[str, Any]]:
    with connect() as conn:
        task = conn.execute("SELECT parent_task_id FROM search_tasks WHERE id=?", (task_id,)).fetchone()
        if not task:
            raise HTTPException(404, "搜索任务不存在")
        task_ids = [task_id]
        if task["parent_task_id"]:
            task_ids.append(int(task["parent_task_id"]))
        placeholders = ",".join("?" for _ in task_ids)
        rows = conn.execute(
            f"""SELECT c.*, p.rank_no, p.original_title, p.year, p.chinese_title,
                      p.tmdb_title,p.tmdb_original_title,p.tmdb_year,p.tmdb_imdb_id,
                      CASE WHEN cart.candidate_id IS NULL THEN 0 ELSE 1 END AS in_cart
               FROM candidates c JOIN playlist_items p ON p.id=c.playlist_item_id
               LEFT JOIN cart_items cart ON cart.candidate_id=c.id
               WHERE c.task_id IN ({placeholders}) ORDER BY p.rank_no, c.ranking""",  # nosec B608
            task_ids,
        ).fetchall()
        site_rows = conn.execute("SELECT name,priority,icon_url FROM pt_sites").fetchall()
    site_profiles = {str(row["name"]).lower(): dict(row) for row in site_rows}
    result = rows_to_dicts(rows)
    for item in result:
        item["context_available"] = item["id"] in raw_candidates
        try:
            item["metadata"] = json.loads(item.pop("metadata_json"))
        except (TypeError, json.JSONDecodeError):
            item["metadata"] = {}
        try:
            item["score_breakdown"] = json.loads(item.get("score_breakdown") or "[]")
        except json.JSONDecodeError:
            item["score_breakdown"] = []
        item["resource_key"] = item.get("resource_key") or resource_fingerprint(item["title"], item.get("size"))
        profile = site_profiles.get(str(item.get("site_name") or "").lower(), {})
        item["site_priority"] = int(profile.get("priority") or 100)
        item["site_icon"] = profile.get("icon_url") or ""
        factor = volume_factor_value(item["metadata"].get("volume_factor"))
        labels = [str(label).lower() for label in item["metadata"].get("labels", [])]
        item["volume_factor"] = factor
        item["is_free"] = factor == 0 or any(label in ("free", "免费", "freeleech") for label in labels)
    groups: dict[tuple[int, str], list[dict[str, Any]]] = {}
    for item in result:
        groups.setdefault((int(item["playlist_item_id"]), item["resource_key"]), []).append(item)
    grouped: list[dict[str, Any]] = []
    for options in groups.values():
        options.sort(key=lambda item: (
            item["site_priority"], 0 if item["is_free"] else 1,
            item["volume_factor"],
            -int(item.get("seeders") or 0), int(item.get("ranking") or 0),
        ))
        primary = dict(options[0])
        primary["site_count"] = len(options)
        primary["site_options"] = [{
            "id": option["id"], "site_name": option.get("site_name"), "seeders": option.get("seeders"),
            "size": option.get("size"), "is_free": option["is_free"], "site_priority": option["site_priority"],
            "volume_factor": option["volume_factor"], "labels": option["metadata"].get("labels", []),
            "in_cart": option.get("in_cart", 0), "context_available": option["context_available"],
        } for option in options]
        factor_label = "免费" if primary["volume_factor"] == 0 else (f"下载 {int(primary['volume_factor'] * 100)}%" if primary["volume_factor"] < 1 else "普通")
        primary["site_selection_reason"] = (
            f"站点优先级 {primary['site_priority']} · {factor_label} · {int(primary.get('seeders') or 0)} 做种"
        )
        primary["in_cart"] = int(any(option.get("in_cart") for option in options))
        grouped.append(primary)
    grouped.sort(key=lambda item: (int(item["rank_no"] or 0), int(item.get("ranking") or 0)))
    return grouped


@app.post("/api/cart/items/{candidate_id}")
async def toggle_cart(candidate_id: str) -> dict[str, Any]:
    with connect() as conn:
        candidate = conn.execute("SELECT eligibility,exclusion_reason FROM candidates WHERE id=?", (candidate_id,)).fetchone()
        if not candidate:
            raise HTTPException(404, "候选不存在")
        exists = conn.execute("SELECT 1 FROM cart_items WHERE candidate_id=?", (candidate_id,)).fetchone()
        if exists:
            conn.execute("DELETE FROM cart_items WHERE candidate_id=?", (candidate_id,))
            return {"candidate_id": candidate_id, "in_cart": False}
        if candidate["eligibility"] != "eligible":
            raise HTTPException(422, f"该资源已被电影策略排除：{candidate['exclusion_reason'] or '不符合允许组合'}")
        if candidate_id not in raw_candidates:
            raise HTTPException(409, "该候选的搜索上下文已失效，请重新搜索后再加入下载列表")
        conn.execute("INSERT INTO cart_items(candidate_id,selected_at) VALUES(?,?)", (candidate_id, utc_now()))
        return {"candidate_id": candidate_id, "in_cart": True}


@app.get("/api/cart")
async def cart() -> list[dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute(
            """SELECT c.id, c.title, c.site_name, c.size, c.resolution, c.library_state, p.original_title
               FROM cart_items cart JOIN candidates c ON c.id=cart.candidate_id
               JOIN playlist_items p ON p.id=c.playlist_item_id ORDER BY cart.selected_at"""
        ).fetchall()
    items = rows_to_dicts(rows)
    for item in items:
        item["context_available"] = item["id"] in raw_candidates
    return items


@app.post("/api/cart/download")
async def download_cart() -> dict[str, Any]:
    with connect() as conn:
        rows = conn.execute(
            """SELECT c.*, p.original_title FROM cart_items cart JOIN candidates c ON c.id=cart.candidate_id
               JOIN playlist_items p ON p.id=c.playlist_item_id ORDER BY cart.selected_at"""
        ).fetchall()
    if not rows:
        raise HTTPException(422, "下载列表为空")
    moviepilot, completed, needs_research, submitted_tasks, expired_items = MoviePilotClient(), 0, 0, [], []
    for candidate in rows:
        raw = raw_candidates.get(candidate["id"])
        if not raw:
            needs_research += 1
            expired_items.append({"candidate_id": candidate["id"], "title": candidate["original_title"]})
            message = "搜索上下文已失效，请重新搜索后加入下载列表"
            with connect() as conn:
                already_recorded = conn.execute(
                    "SELECT 1 FROM download_history WHERE candidate_id=? AND success=0 AND message=? LIMIT 1",
                    (candidate["id"], message),
                ).fetchone()
                if not already_recorded:
                    conn.execute(
                        """INSERT INTO download_history(
                               candidate_id,playlist_item_id,title,torrent_name,site_name,success,message,created_at
                           ) VALUES(?,?,?,?,?,?,?,?)""",
                        (candidate["id"], candidate["playlist_item_id"], candidate["original_title"], candidate["title"], candidate["site_name"], 0, message, utc_now()),
                    )
            continue
        try:
            if not raw.get("media"):
                raise RuntimeError("缺少媒体信息，无法应用 MoviePilot 分类规则")
            # 固定走 MoviePilot DownloadChain：它补全 TMDB 媒体信息、按 MP 分类目录选择路径，
            # 再交由 Transmission 写入 MOVIEPILOT 与站点标签，供 MP 后续整理。
            response = await moviepilot.download(raw["media"], raw["torrent"], downloader="Transmission")
            success = bool(response.get("success", True)) if isinstance(response, dict) else True
            message = response.get("message") or response.get("hash") if isinstance(response, dict) else None
            message = sanitize_sensitive_text(message) if message else None
            submission_hash = str(response.get("hash") or "").strip() or None if isinstance(response, dict) else None
            if isinstance(response, dict):
                submitted_tasks.append({"candidate_id": candidate["id"], "hash": response.get("hash"), "mode": "moviepilot"})
        except Exception as exc:
            success, message, submission_hash = False, safe_error(exc), None
        with connect() as conn:
            conn.execute(
                """INSERT INTO download_history(
                       candidate_id,playlist_item_id,title,torrent_name,site_name,submission_hash,success,message,created_at
                   ) VALUES(?,?,?,?,?,?,?,?,?)""",
                (candidate["id"], candidate["playlist_item_id"], candidate["original_title"], candidate["title"], candidate["site_name"], submission_hash, int(success), message, utc_now()),
            )
            if success:
                conn.execute("DELETE FROM cart_items WHERE candidate_id=?", (candidate["id"],))
                completed += 1
    if needs_research and completed == 0:
        raise HTTPException(409, f"下载列表中 {needs_research} 个资源的搜索上下文已失效，请重新搜索后加入下载列表")
    return {
        "submitted": completed, "needs_research": needs_research, "expired_items": expired_items,
        "mode": "moviepilot", "tasks": submitted_tasks,
    }


@app.get("/api/history")
async def history() -> list[dict[str, Any]]:
    return await projected_download_history()


@app.get("/")
async def index() -> FileResponse:
    return FileResponse(Path(__file__).parent / "static" / "index.html")
